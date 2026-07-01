"""web/backend/files.py — 업로드 파일(Excel/CSV/PDF) 텍스트 추출.

LLM이 파일 내용을 '분석'하고 그에 대해 질문할 수 있도록, 온프렘에서 파일을 텍스트로 추출한다.
외부 전송 없이 서버 메모리에서만 처리한다(데모: 세션 종료 시 폐기).
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field

_MAX_CHARS = 12000  # LLM 컨텍스트 보호용 상한


@dataclass
class ExtractedFile:
    name: str
    kind: str            # excel | csv | pdf | unknown
    text: str            # 추출 텍스트(상한 적용)
    summary: str         # 짧은 요약(행/열/시트/페이지 수 등)
    truncated: bool = False
    meta: dict = field(default_factory=dict)


def _clip(text: str) -> tuple[str, bool]:
    if len(text) <= _MAX_CHARS:
        return text, False
    return text[:_MAX_CHARS] + "\n…(이하 생략)", True


def extract(filename: str, data: bytes) -> ExtractedFile:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm", ".xls")):
        return _excel(filename, data)
    if lower.endswith(".csv"):
        return _csv(filename, data)
    if lower.endswith(".pdf"):
        return _pdf(filename, data)
    text, tr = _clip(_decode(data))
    return ExtractedFile(filename, "unknown", text, "형식 미상 텍스트로 읽음", tr)


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "cp949", "euc-kr", "latin-1"):
        try:
            return data.decode(enc)
        except Exception:  # noqa: BLE001
            continue
    return data.decode("utf-8", errors="replace")


def _excel(filename: str, data: bytes) -> ExtractedFile:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    total_rows = 0
    try:
        sheet_count = len(wb.worksheets)
        for ws in wb.worksheets:
            parts.append(f"[시트: {ws.title}]")
            rows = 0
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if not any(c.strip() for c in cells):
                    continue
                parts.append(" | ".join(cells))
                rows += 1
                total_rows += 1
                if rows >= 200:  # 시트당 상한
                    parts.append("…(행 생략)")
                    break
    finally:
        wb.close()
    text, tr = _clip("\n".join(parts))
    summary = f"Excel · 시트 {sheet_count}개 · 데이터행 약 {total_rows}행"
    return ExtractedFile(filename, "excel", text, summary, tr,
                         {"sheets": sheet_count, "rows": total_rows})


def _csv(filename: str, data: bytes) -> ExtractedFile:
    text_raw = _decode(data)
    reader = csv.reader(io.StringIO(text_raw))
    lines = []
    n = 0
    for row in reader:
        lines.append(" | ".join(row))
        n += 1
        if n >= 300:
            lines.append("…(행 생략)")
            break
    text, tr = _clip("\n".join(lines))
    return ExtractedFile(filename, "csv", text, f"CSV · 약 {n}행", tr, {"rows": n})


def _pdf(filename: str, data: bytes) -> ExtractedFile:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    parts = []
    try:
        pages = doc.page_count
        for i, page in enumerate(doc):
            parts.append(f"[p.{i + 1}]")
            parts.append(page.get_text("text"))
            if i >= 30:  # 페이지 상한
                parts.append("…(이후 페이지 생략)")
                break
    finally:
        doc.close()
    body = "\n".join(parts).strip()
    if not body:
        body = "(추출된 텍스트가 없습니다 — 스캔 PDF일 수 있어 OCR이 필요합니다.)"
    text, tr = _clip(body)
    return ExtractedFile(filename, "pdf", text, f"PDF · {pages}페이지", tr, {"pages": pages})
