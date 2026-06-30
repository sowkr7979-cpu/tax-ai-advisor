"""tests/test_tax_plan_report.py — 시나리오 Tax Plan 렌더러 검증.

보장:
  1. 재무제표 쟁점 행 강조(음영) + 근거 하이퍼링크.
  2. 영어/개발어 0(회계사 자연어), Noto Sans KR.
  3. 인용 날조 차단(source_objects 없는 인용은 fail-closed).
  4. 그림·도표(차트) 임베드(matplotlib).

모두 오프라인(법령 fixture replay) — 실키/네트워크 불필요.
"""
from __future__ import annotations

import dataclasses
import re

import pytest
from docx import Document
from docx.oxml.ns import qn

from src.tax_plan import build_demo_tax_plan
from src.tax_plan_report import TaxPlanValidationError, build_tax_plan_docx


def _visible(path):
    d = Document(str(path))
    chunks = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                chunks.append(c.text)
    return "\n".join(chunks)


def _shaded_cells(path):
    d = Document(str(path))
    n = 0
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                tcpr = c._tc.find(qn("w:tcPr"))
                if tcpr is not None and tcpr.find(qn("w:shd")) is not None:
                    n += 1
    return n


def test_tax_plan_structure_highlights_links_korean(tmp_path):
    data, titles, articles = build_demo_tax_plan()
    out = build_tax_plan_docx(data, tmp_path / "plan.docx", titles=titles, articles=articles,
                              charts_dir=tmp_path / "ch", with_charts=False)
    body = _visible(out)
    d = Document(str(out))

    # (2) 폰트 + 영어 0
    assert d.styles["Normal"].font.name == "Noto Sans KR"
    assert re.findall(r"[A-Za-z]+", body) == [], "회계사 보고서 본문에 영어 노출"

    # (1) 근거 하이퍼링크 — 5개 법령, 국가법령정보센터
    targets = [r.target_ref for r in d.part.rels.values() if "hyperlink" in r.reltype]
    assert len(targets) >= 5 and all("law.go.kr" in t for t in targets)

    # 재무제표 쟁점 강조(음영 셀 존재)
    assert _shaded_cells(out) > 10

    # 핵심 장(章) + 수치 존재
    for s in ("재무상태표", "손익계산서", "잉여금 환원 시나리오 비교", "가업승계 증여세",
              "관련 법령", "실행 타임라인"):
        assert s in body, f"섹션 '{s}' 누락"
    for num in ("745", "110", "25", "200", "22"):
        assert num in body


def test_tax_plan_rejects_unverified_citation(tmp_path):
    """인용에 대응하는 source_objects(등록 소스객체)가 없으면 날조 차단 — fail-closed."""
    data, titles, articles = build_demo_tax_plan()
    broken = dataclasses.replace(data, source_objects=[])
    with pytest.raises(TaxPlanValidationError):
        build_tax_plan_docx(broken, tmp_path / "x.docx", titles=titles, articles=articles,
                            with_charts=False)


def test_tax_plan_embeds_charts(tmp_path):
    """그림·도표(지분구조도·타임라인·세부담 비교·증여세 워터폴) 4종이 임베드된다(matplotlib)."""
    data, titles, articles = build_demo_tax_plan()
    out = build_tax_plan_docx(data, tmp_path / "plan.docx", titles=titles, articles=articles,
                              charts_dir=tmp_path / "ch", with_charts=True)
    d = Document(str(out))
    images = [r for r in d.part.rels.values() if "image" in r.reltype]
    assert len(images) == 4
