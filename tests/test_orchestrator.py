"""tests/test_orchestrator.py — 오케스트레이터 end-to-end replay 스모크 테스트.

The orchestrator chains the 6 slices into "새 질문 → 검토패키지 DOCX". These tests
prove the REPLAY path is hermetic + DETERMINISTIC (no network/key) and that the
core invariants hold:

  * §4-1 chain produces an 11목차 package: 4 version-object 인용 · 보수/중립/적극 3 선택지 ·
    HALU-008 검토항목(고위험 검토경고 포함);
  * 3소스 종합(①②③)은 AGREE + lineage 무결성(추적불가 0 · 신규 인용 0);
  * 인용 검증(SourceRegistry) — validate_draft_package 통과(날조 차단);
  * 테넌트 격리키(client_id) 전파;
  * HITL(OUT-004) — 고객 전달본은 승인 proof 없이는 생성 차단(내부본 fallback), proof 있으면 생성;
  * 결정성 — 동일 입력 2회 → 동일 패키지/DOCX 내용.

Hermetic: everything REPLAYS from tests/fixtures/ (LLM gen+strategy, model2vec 임베딩,
official web content, 법령 본문). A missing fixture is fail-closed (the run raises).
"""

from __future__ import annotations

import json

import pytest

from docx import Document

from src.draft import REQUIRED_SECTIONS, validate_draft_package
from src.orchestrator import (
    DEFAULT_COMPANY_FIXTURE,
    CompanyProfile,
    Orchestrator,
)

_REQUIRED_OPTION_KEYS = {"보수", "중립", "적극"}


def _company() -> CompanyProfile:
    return CompanyProfile.from_fixture(DEFAULT_COMPANY_FIXTURE)


def _run(out_path=None, audience="internal", approve_demo=False):
    company = _company()
    orch = Orchestrator(mode="replay")
    return orch.run(
        company=company, question=company.default_question, out_path=out_path,
        audience=audience, approve_demo=approve_demo, write_docx=out_path is not None,
    )


def _docx_content(path) -> list[str]:
    d = Document(str(path))
    parts = [f"P:{p.text}" for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            parts.append("T:" + "|".join(c.text for c in r.cells))
    return parts


# --------------------------------------------------------------------------- #
# Package shape (§4-1 산출)
# --------------------------------------------------------------------------- #
def test_orchestrator_builds_11section_package():
    result = _run()
    pkg = result.package
    # 11목차 — validate enforces every required section + 무인용 단정 금지(SourceRegistry).
    validate_draft_package(pkg)  # must not raise (인용 검증 통과 = 날조 차단)
    assert len(REQUIRED_SECTIONS) == 11
    # 4 version-object 인용 (제25/24/27의2/28조)
    assert len(pkg.citations) == 4
    # 보수/중립/적극 3 선택지
    assert {o.option_key for o in pkg.strategy_options} == _REQUIRED_OPTION_KEYS
    # 검토항목(HALU-008) — 고위험 검토경고 포함
    from src.review_items import has_high_risk_warning
    assert pkg.review_items and has_high_risk_warning(pkg.review_items)


def test_three_source_synthesis_agree_lineage_intact():
    result = _run()
    syn = result.synthesis
    assert syn is not None
    # 3소스(①②③) 모두 응답 → 합의
    answered = [c for c in result.contributions if c.answered]
    assert len(answered) == 3
    assert not syn.abstained
    # lineage 무결성: 추적불가 0 · 신규 인용 0 (ORCH-012/HALU-014)
    assert syn.untraceable_claims() == []
    assert syn.new_citations == []


def test_strategy_options_cite_only_retrieved_version_objects():
    result = _run()
    strat = result.strategy
    assert strat is not None and strat.has_all_three
    cit_ids = {c.citation_id for c in result.package.citations}
    # 모든 선택지 인용은 회수된 버전 객체(패키지 인용)에 한정 — 날조 id 없음
    for opt in strat.options:
        assert opt.citation_ids
        assert all(cid in cit_ids for cid in opt.citation_ids)


def test_tenant_isolation_key_propagated():
    result = _run()
    company = _company()
    assert result.package.client_id == company.client_id
    # RAG 채널(②)의 SourceAnswer 가 동일 격리키를 운반
    rag = next((c for c in result.contributions if c.channel_label == "②"), None)
    assert rag is not None
    if rag.answered:
        assert rag.source_answer.client_id == company.client_id


# --------------------------------------------------------------------------- #
# 결정성 (replay)
# --------------------------------------------------------------------------- #
def test_replay_is_deterministic_package(tmp_path):
    a = _run(out_path=tmp_path / "a.docx")
    b = _run(out_path=tmp_path / "b.docx")
    # 패키지 직렬화(JSON)가 동일
    fa = json.dumps(a.to_fixture(), ensure_ascii=False, sort_keys=True)
    fb = json.dumps(b.to_fixture(), ensure_ascii=False, sort_keys=True)
    assert fa == fb
    # DOCX 추출 내용(문단·표)이 동일 (바이트 차이는 docx zip 타임스탬프 뿐)
    assert _docx_content(tmp_path / "a.docx") == _docx_content(tmp_path / "b.docx")


def test_docx_has_all_11_sections(tmp_path):
    out = tmp_path / "검토패키지.docx"
    result = _run(out_path=out)
    assert result.docx_path is not None and out.exists()
    content = _docx_content(out)
    text = "\n".join(content)
    for section in REQUIRED_SECTIONS:
        assert section in text, f"누락 섹션: {section}"
    # 선택지 비교표(보수/중립/적극)가 표로 렌더됨
    assert any("세무 리스크(과세논리)" in p for p in content)


# --------------------------------------------------------------------------- #
# HITL 불변식 (OUT-004)
# --------------------------------------------------------------------------- #
def test_client_deliverable_blocked_without_approval(tmp_path):
    out = tmp_path / "client.docx"
    result = _run(out_path=out, audience="client", approve_demo=False)
    # 승인 proof 없음 → 고객 전달본 차단, 내부본만 생성(fail-closed)
    assert result.is_client_deliverable is False
    assert out.exists()
    assert any("고객 전달본 차단" in line for line in result.log)


def test_client_deliverable_with_demo_approval_excludes_internal(tmp_path):
    out = tmp_path / "client.docx"
    result = _run(out_path=out, audience="client", approve_demo=True)
    assert result.is_client_deliverable is True
    content = "\n".join(_docx_content(out))
    # OUT-004: 고객 전달본은 내부 전용 섹션(7. 쟁점별 검토 메모 / 10. 회계사 검토 필요사항) 제외
    assert "7. 쟁점별 검토 메모" not in content
    assert "10. 회계사 검토 필요사항" not in content
    # 그래도 선택지 비교표·법령 근거는 포함
    assert "6. 선택지별 세부담·리스크 비교표" in content


# --------------------------------------------------------------------------- #
# fail-closed 하드닝 (codex P0 회귀 — 우회/은폐 방지)
# --------------------------------------------------------------------------- #
def test_out003_validation_runs_even_without_docx(monkeypatch):
    """OUT-003 인용 검증은 DOCX 출력 여부와 무관하게 항상 실행된다(write_docx=False 우회 차단)."""
    import src.orchestrator as orch_mod

    calls = {"n": 0}
    real = orch_mod.validate_draft_package

    def spy(pkg):
        calls["n"] += 1
        return real(pkg)

    monkeypatch.setattr(orch_mod, "validate_draft_package", spy)
    _run(out_path=None)  # out 미지정 → write_docx=False
    assert calls["n"] == 1, "패키지만 받아가는 경로에서도 OUT-003 게이트가 실행되어야 함"


def _patched_rag(monkeypatch, exc: Exception):
    """내부 RAG 채널의 answer 가 exc 를 던지도록 교체."""
    import src.orchestrator as orch_mod

    class _BoomIndex:
        def __init__(self, *a, **k):
            pass

        def ingest(self, *a, **k):
            pass

        def answer(self, *a, **k):
            raise exc

    monkeypatch.setattr(orch_mod, "InternalRagIndex", _BoomIndex)


def test_replay_failclosed_on_isolation_error(monkeypatch):
    """격리(IsolationError) 위반은 어떤 모드에서도 절대 삼키지 않고 즉시 전파(TENANT_LEAK=0)."""
    from src.isolation import IsolationError

    _patched_rag(monkeypatch, IsolationError("simulated tenant breach"))
    with pytest.raises(IsolationError):
        _run()  # graceful-degrade 로 은폐되지 않고 fail-closed


def test_replay_failclosed_on_fixture_failure(monkeypatch):
    """replay 는 hermetic — 연구 채널의 누락/재현불가 실패는 silent degrade 가 아니라 전파."""
    _patched_rag(monkeypatch, RuntimeError("simulated missing/mismatched fixture"))
    with pytest.raises(RuntimeError):
        _run()  # mode=replay → fail-closed (Invariant 5)


# --------------------------------------------------------------------------- #
# 비-데모 입력에 잘못된 법리 요약 금지 (codex stop-time 회귀)
# --------------------------------------------------------------------------- #
def test_summary_derives_from_primary_issue_not_hardcoded_meal():
    """주쟁점이 기업업무추진비가 아니면 요약/결론/권고에 접대비 서사를 출력하지 않는다."""
    from src.orchestrator import Orchestrator

    company = _company()
    # 비-데모 주쟁점(지급이자, 제28조)
    primary = {
        "issue_key": "지급이자", "article": "제28조",
        "article_title": "지급이자의 손금불산입",
        "title": "지급이자 손금불산입(가지급금 등)",
    }
    # 지급이자 주쟁점에서는 웹이 보류(①②만 응답) — 거짓 웹 회수 주장이 없어야 함
    exec_summary, conclusion, order = Orchestrator._summary_texts(
        company, primary, {"①", "②"})
    blob = exec_summary + " " + conclusion + " " + " ".join(order)
    # 접대비/승용차 전용 서사가 새지 않음
    for meal_token in ("기업업무추진비", "접대비", "적격증빙", "예규·심판례", "운행기록부"):
        assert meal_token not in blob, f"비-접대비 주쟁점인데 접대비/승용차 서사 누출: {meal_token!r}"
    # 실제 주쟁점이 반영됨
    assert "지급이자" in exec_summary and "제28조" in exec_summary
    assert "지급이자" in conclusion
    # 웹 미응답을 정직하게 기술(③웹 회수 주장 금지) — 회수 주장부(권위 위계 앞)에 웹 없음,
    # 단 정직한 보류 고지는 포함.
    retrieved_claim = exec_summary.split("권위 위계")[0]
    assert "③공식웹" not in retrieved_claim and "③웹" not in retrieved_claim, "웹 보류인데 ③웹 회수 주장 누출"
    assert "③공식웹은 적용 가능한 공식근거가 없어 보류" in exec_summary


def test_summary_for_meal_issue_keeps_specialized_clauses():
    """주쟁점이 실제 기업업무추진비면 예규·심판례 심화 절을 포함한다(데모 경로)."""
    from src.orchestrator import Orchestrator

    company = _company()
    primary = {
        "issue_key": "기업업무추진비", "article": "제25조",
        "article_title": "기업업무추진비의 손금불산입",
        "title": "기업업무추진비(접대비) 한도·적격증빙",
    }
    # 데모 경로: 3소스 모두 응답
    exec_summary, conclusion, order = Orchestrator._summary_texts(
        company, primary, {"①", "②", "③"})
    assert "예규·심판례" in exec_summary
    assert any("예규·심판례" in s for s in order)
    assert "기업업무추진비(접대비) 한도·적격증빙" in exec_summary
    # 웹 응답 → ③공식웹 회수 명시, 보류 문구 없음
    assert "③공식웹" in exec_summary and "보류" not in exec_summary


def test_evidence_requests_match_detected_issues():
    """추가 요청 자료는 식별된 쟁점에만 해당 — 무관 쟁점의 접대비/승용차 자료를 요구하지 않음."""
    from src.orchestrator import Orchestrator

    orch = Orchestrator(mode="replay")
    issues = [{"issue_key": "지급이자", "article": "제28조",
               "article_title": "지급이자의 손금불산입", "title": "지급이자 손금불산입(가지급금 등)"}]
    reqs = orch._evidence_requests(issues, [])
    blob = " ".join(reqs)
    assert "지급이자" in blob and "가지급금" in blob
    assert "접대비" not in blob and "운행기록부" not in blob and "기부금" not in blob
    assert any("이사회의사록" in r for r in reqs)  # 공통 거버넌스 자료


def test_spot_issues_promotes_question_named_issue():
    """질문이 특정 쟁점을 명시하면 그 쟁점을 주쟁점으로(질문 무관 고정 차단)."""
    from src.orchestrator import Orchestrator

    orch = Orchestrator(mode="replay")
    company = _company()  # TB 에 기업업무추진비 + 지급이자 등 포함
    # 질문이 지급이자를 명시 → 주쟁점이 기업업무추진비가 아닌 지급이자
    issues = orch._spot_issues(company, "지급이자 가지급금 인정이자 손금불산입 검토해줘")
    assert issues[0]["issue_key"] == "지급이자"
    # 질문이 기업업무추진비/접대비를 명시(데모 기본) → 주쟁점 기업업무추진비(기존 동작 유지)
    issues2 = orch._spot_issues(company, company.default_question)
    assert issues2[0]["issue_key"] == "기업업무추진비"


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(pytest.main([__file__, "-v"]))
