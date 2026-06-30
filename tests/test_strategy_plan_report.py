"""tests/test_strategy_plan_report.py — 다(多)전략 절세 종합검토 렌더러 검증.

보장:
  1. 재무제표를 DART 공시 그대로(원 단위 raw 값)로 첨부 + 쟁점 행 강조 + 근거 링크.
  2. 여러 전략을 재무제표에 대해 스크리닝(적용/조건부/미해당) — '미해당'도 사유와 함께 노출.
  3. 회계사 자연어(영어/개발어 0), Noto Sans KR.
  4. 인용 날조 차단(source_objects 없는 인용은 fail-closed).
  5. 그림·도표(스크리닝 맵·기대효과 등급·실행 로드맵) 임베드.

모두 오프라인(법령 fixture 재현) — 실키/네트워크 불필요.
"""
from __future__ import annotations

import dataclasses
import re

import pytest
from docx import Document
from docx.oxml.ns import qn

from src.strategy_plan import build_demo_strategy_plan
from src.strategy_plan_report import StrategyPlanValidationError, build_strategy_plan_docx


def _visible(path):
    d = Document(str(path))
    chunks = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                chunks.append(c.text)
    return "\n".join(chunks)


def _shaded_cells(path):
    d = Document(str(path))
    n = 0
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                tcpr = c._tc.find(qn("w:tcPr"))
                if tcpr is not None and tcpr.find(qn("w:shd")) is not None:
                    n += 1
    return n


def test_strategy_plan_structure_links_korean_raw_won(tmp_path):
    data, titles, articles = build_demo_strategy_plan()
    out = build_strategy_plan_docx(data, tmp_path / "plan.docx", titles=titles, articles=articles,
                                   charts_dir=tmp_path / "ch", with_charts=False)
    body = _visible(out)
    d = Document(str(out))

    # (3) 폰트 + 영어 0
    assert d.styles["Normal"].font.name == "Noto Sans KR"
    assert re.findall(r"[A-Za-z]+", body) == [], "회계사 보고서 본문에 영어 노출"

    # (1) 재무제표 = DART 원본(원 단위 raw 값) — 억원 환산이 아니라 실값
    assert "611,759,633,069" in body, "이익잉여금 원 단위 실값 누락"
    assert "278,186,189,427" in body, "매출액 원 단위 실값 누락"
    assert "단위: 원" in body

    # (1) 근거 하이퍼링크 — 다수 법령, 국가법령정보센터
    targets = [r.target_ref for r in d.part.rels.values() if "hyperlink" in r.reltype]
    assert len(targets) >= 8 and all("law.go.kr" in t for t in targets)

    # 쟁점 강조(음영 셀 다수)
    assert _shaded_cells(out) > 20

    # 핵심 장(章) 존재 — 재무제표 + 계정별원장·분개장 + 스크리닝 + 상세(필요 조치·자료·사후관리)
    for s in ("재무상태표", "손익계산서", "계정별원장", "분개장", "절세전략 종합 스크리닝",
              "적용 권고 전략", "필요 조치", "필요 자료", "사후관리", "실행 우선순위", "관련 법령"):
        assert s in body, f"섹션/항목 '{s}' 누락"
    # 분개장의 원 단위 거래 금액(배당 300억 등)이 그대로 노출
    assert "30,000,000,000" in body, "분개장 거래 금액(원) 누락"


def test_strategy_plan_screens_many_with_not_applicable(tmp_path):
    """가업승계 하나가 아니라 다수 전략을 스크리닝하고, 미해당도 사유와 함께 노출한다."""
    data, titles, articles = build_demo_strategy_plan()
    assert len(data.strategies) >= 12, "전략 카탈로그가 충분히 많아야 함"
    assert sum(1 for s in data.strategies if s.applies) >= 8, "적용 전략 다수"
    assert any(not s.applies for s in data.strategies), "미해당 전략도 있어야(스크리닝 투명성)"

    out = build_strategy_plan_docx(data, tmp_path / "plan.docx", titles=titles, articles=articles,
                                   charts_dir=tmp_path / "ch", with_charts=False)
    body = _visible(out)
    # 미해당으로 스크리닝된 중소기업 특별세액감면이 사유와 함께 노출
    assert "미해당" in body
    assert "중소기업" in body
    # 다양한 카테고리가 등장
    for cat in ("잉여금 환원·자본거래", "세액공제·감면", "손금·익금 최적화"):
        assert cat in body, f"카테고리 '{cat}' 누락"


def test_strategy_plan_rejects_unverified_citation(tmp_path):
    """인용에 대응하는 source_objects가 없으면 날조 차단 — fail-closed."""
    data, titles, articles = build_demo_strategy_plan()
    broken = dataclasses.replace(data, source_objects=[])
    with pytest.raises(StrategyPlanValidationError):
        build_strategy_plan_docx(broken, tmp_path / "x.docx", titles=titles, articles=articles,
                                 with_charts=False)


def test_strategy_plan_three_source_research(tmp_path):
    """§2 — 법령①·내부자료②·공식웹③ 3소스 답변 + 종합의견이 모두 노출된다."""
    data, titles, articles = build_demo_strategy_plan()
    assert data.research is not None, "3소스 리서치 데이터 필요"
    out = build_strategy_plan_docx(data, tmp_path / "plan.docx", titles=titles, articles=articles,
                                   charts_dir=tmp_path / "ch", with_charts=False)
    body = _visible(out)

    # 3개 채널 + 종합의견 + 검토질문이 모두 본문에 등장
    for token in ("3소스 교차검증", "①법령", "②내부자료", "③공식웹", "종합의견", "검토 질문"):
        assert token in body, f"리서치 섹션 항목 '{token}' 누락"
    # 채널 법령 인용(제18조의2)이 §2 에서 하이퍼링크로 노출(전체 링크 수 증가)
    d = Document(str(out))
    targets = [r.target_ref for r in d.part.rels.values() if "hyperlink" in r.reltype]
    assert any("law.go.kr" in t for t in targets)
    # 영어 0 유지(오프라인 데모 — RAG 회수 근거는 비어 있음)
    assert re.findall(r"[A-Za-z]+", body) == [], "리서치 섹션에 영어 노출"


def test_strategy_plan_rag_research_attach_degrades_when_no_index(tmp_path, monkeypatch):
    """attach_rag_db_research 는 RAG 인덱스가 없으면 ②채널을 보류(침묵)로 표시(날조 0)."""
    from src import strategy_plan as sp

    data, _, _ = build_demo_strategy_plan()

    def _boom(*a, **k):
        raise FileNotFoundError("no index (test)")

    monkeypatch.setattr(sp, "load_index", _boom, raising=False)
    # load_index 는 함수 내부에서 import 되므로 모듈 경로를 직접 패치
    import src.rag_db_index as ragdb
    monkeypatch.setattr(ragdb, "load_index", _boom, raising=False)

    out = sp.attach_rag_db_research(data)
    rag_ch = [c for c in out.research.channels if c.authority == "실무서"][0]
    assert rag_ch.abstained is True and not rag_ch.rag_evidence


def test_strategy_plan_embeds_charts(tmp_path):
    """그림·도표(스크리닝 맵·기대효과 등급·실행 로드맵) 3종이 임베드된다(matplotlib)."""
    data, titles, articles = build_demo_strategy_plan()
    out = build_strategy_plan_docx(data, tmp_path / "plan.docx", titles=titles, articles=articles,
                                   charts_dir=tmp_path / "ch", with_charts=True)
    d = Document(str(out))
    images = [r for r in d.part.rels.values() if "image" in r.reltype]
    assert len(images) == 3
