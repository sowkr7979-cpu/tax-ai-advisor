"""tests/test_scenario_planner.py — 경우의 수 시나리오 모델·플로우차트·보고서 검증.

데모 시나리오의 구조 무결성(분기·경우의 수·근거·위험등급), fail-closed 검증, 플로우차트/
매트릭스 PNG 렌더, DOCX 산출(표·이미지·분개장/원장 부록)을 확인한다. 외부 키·RAG 인덱스
없이 통과한다(with_rag=False / OCR 무관).
"""
from __future__ import annotations

import dataclasses

import pytest

import src.scenario_planner as sp


def test_demo_scenarios_structure():
    scns = sp.demo_scenarios()
    assert len(scns) == 3
    keys = {s.key for s in scns}
    assert keys == {"TREASURY", "EXECPAY", "LOAN"}
    for s in scns:
        s.validate()                      # 구조 무결성(예외 없음)
        assert len(s.gates) >= 2
        assert len(s.leaves) >= 2
        assert s.recommended_case in {lf.case_label for lf in s.leaves}
        for lf in s.leaves:
            assert lf.risk_level in sp.RISK_LEVELS
            assert lf.basis                # 근거 필수(날조 방지)
        assert s.citations                 # 시나리오 근거 법령 필수
        # 사후관리·필요조치·실행계획(요구사항 3)
        assert s.aftercare and s.actions and s.plan


def test_recommended_case_is_lowest_risk():
    # 권고 경로는 '안전' 등급이어야 한다(데모 설계 불변식).
    for s in sp.demo_scenarios():
        rec = next(lf for lf in s.leaves if lf.case_label == s.recommended_case)
        assert rec.risk_level == "안전"


def _base() -> sp.Scenario:
    return sp.demo_scenarios()[0]


def test_validate_rejects_too_few_gates():
    s = dataclasses.replace(_base(), gates=_base().gates[:1])
    with pytest.raises(ValueError):
        s.validate()


def test_validate_rejects_bad_risk_level():
    s = _base()
    bad = dataclasses.replace(s.leaves[0], risk_level="치명적")
    s2 = dataclasses.replace(s, leaves=(bad, *s.leaves[1:]))
    with pytest.raises(ValueError):
        s2.validate()


def test_validate_rejects_empty_basis():
    s = _base()
    bad = dataclasses.replace(s.leaves[0], basis=())
    s2 = dataclasses.replace(s, leaves=(bad, *s.leaves[1:]))
    with pytest.raises(ValueError):
        s2.validate()


def test_validate_rejects_whitespace_basis():
    # 빈 문자열/공백 근거로 날조방지 가드를 우회하지 못한다(codex 적발).
    s = _base()
    bad = dataclasses.replace(s.leaves[0], basis=("  ",))
    s2 = dataclasses.replace(s, leaves=(bad, *s.leaves[1:]))
    with pytest.raises(ValueError):
        s2.validate()
    # 시나리오 근거(citations)도 동일
    s3 = dataclasses.replace(_base(), citations=("",))
    with pytest.raises(ValueError):
        s3.validate()


def test_validate_rejects_recommended_not_in_leaves():
    s = dataclasses.replace(_base(), recommended_case="경우 ⑨")
    with pytest.raises(ValueError):
        s.validate()


def test_flowchart_and_matrix_render(tmp_path):
    import src.scenario_flowchart as sfc

    s = _base()
    f = sfc.save_scenario_flowchart(s, tmp_path / "flow.png")
    m = sfc.save_case_matrix(s, tmp_path / "matrix.png")
    assert f.exists() and f.stat().st_size > 5000
    assert m.exists() and m.stat().st_size > 5000


def test_law_search_url_is_portal_search():
    from src.scenario_report import _law_search_url

    url = _law_search_url("법인세법 제52조(부당행위계산의 부인)")
    assert url.startswith("https://www.law.go.kr/")
    assert "query=" in url
    # 법령명만 검색(괄호 설명 제외 — 임의 조문ID 합성 안 함)
    assert "%EB" in url  # url-encoded 한글


def test_build_scenario_docx(tmp_path):
    from docx import Document

    from src.scenario_report import build_scenario_docx

    out = build_scenario_docx(
        sp.demo_scenarios(), tmp_path / "scn.docx",
        with_rag=False, with_charts=True, charts_dir=tmp_path / "_charts",
    )
    assert out.exists()
    d = Document(str(out))
    # 플로우차트3 + 매트릭스3 = 인라인 이미지 6
    assert len(d.inline_shapes) == 6
    heads = [p.text for p in d.paragraphs
             if p.style.name.startswith("Heading") or p.style.name == "Title"]
    # 방법론 + 시나리오3 + 분개장/원장 부록 + 종합유의
    assert any("방법론" in h for h in heads)
    assert any("분개장" in h for h in heads)
    assert any("원장" in h for h in heads)
    # 사후관리·필요조치·실행계획 섹션 존재
    assert any("사후관리" in h and "실행계획" in h for h in heads)


def test_build_scenario_docx_no_charts(tmp_path):
    from docx import Document

    from src.scenario_report import build_scenario_docx

    out = build_scenario_docx(
        sp.demo_scenarios(), tmp_path / "scn2.docx",
        with_rag=False, with_charts=False,
    )
    d = Document(str(out))
    assert len(d.inline_shapes) == 0       # 그림 생략
    assert len(d.tables) >= 3              # 표(경우의 수·실행계획·분개 등)는 유지
