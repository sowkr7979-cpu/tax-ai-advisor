"""tests/test_dart_case.py — DART 실재무 그라운딩 종합세무검토·의사결정 보고서 검증.

네트워크 없이 합성 ``CompanyFinancials`` 로 회사개요·종합세무검토·회사 시나리오([사실관계])·
DOCX 산출(회사개요 표·종합세무검토 표·시나리오4·분개장/원장/증빙 부록)을 확인한다.
"""
from __future__ import annotations

import src.dart_case_data as dc
from src.dart_fetch import CompanyFinancials, FsLine


def _fake_fin() -> CompanyFinancials:
    lines = [
        FsLine("자산총계", 7.1e11, "BS"),
        FsLine("자본총계", 5.4e11, "BS"),
        FsLine("자본금", 1.27e10, "BS"),
        FsLine("이익잉여금", 6.1e11, "BS"),
        FsLine("매출액", 5.59e11, "IS"),
        FsLine("영업이익", 2.55e11, "IS"),
        FsLine("당기순이익", 1.53e11, "IS"),
    ]
    return CompanyFinancials(corp_name="테스트반도체", corp_code="00000000",
                             stock_code="000000", year=2024, fs_div="CFS",
                             lines=tuple(lines))


def test_company_overview_uses_real_figures():
    rows, note, src_note = dc.company_overview(_fake_fin())
    d = dict(rows)
    assert "자산총계" in d and "이익잉여금" in d
    assert "6,100억" in d["이익잉여금"] or "6,100" in d["이익잉여금"]
    # 자본금 대비 이익잉여금 배수가 note 에 반영(약 48배)
    assert "배" in note
    assert "가상" in src_note and "DART" in src_note  # 데이터 출처·dummy 고지


def test_comprehensive_review_rows():
    rows = dc.comprehensive_review(_fake_fin())
    assert len(rows) >= 5
    for acct, issue, direction, basis in rows:
        assert acct and issue and direction and basis.strip()  # 근거 비어있지 않음


def test_company_scenarios_valid_and_grounded():
    scns = dc._company_scenarios(_fake_fin())
    assert len(scns) == 4
    assert {s.key for s in scns} == {"TREASURY", "EXECPAY", "LOAN", "RND"}
    for s in scns:
        s.validate()                       # 구조 무결성(근거 필수 등)
        assert s.facts.strip()             # [사실관계] 합성됨
        assert "테스트반도체" in s.facts    # 회사 그라운딩


def test_rnd_scenario_is_tax_saving_positive():
    rnd = dc._rnd_scenario(_fake_fin())
    rnd.validate()
    # 권고 경로가 '안전'(공제 극대화) + 조특법 근거
    rec = next(lf for lf in rnd.leaves if lf.case_label == rnd.recommended_case)
    assert rec.risk_level == "안전"
    assert any("조세특례제한법" in c or "조특법" in c for c in rnd.citations)


def test_build_company_case_docx(tmp_path):
    from docx import Document

    from src.scenario_report import build_company_case_docx

    fin = _fake_fin()
    rows, note, src_note = dc.company_overview(fin)
    scns = dc._company_scenarios(fin)
    out = build_company_case_docx(
        scns, tmp_path / "case.docx",
        company="테스트반도체(주)", as_of="2026년 7월 1일",
        overview_rows=rows, overview_note=note,
        tax_review_rows=dc.comprehensive_review(fin),
        evidence=dc.evidence_list(), data_source_note=src_note,
        with_rag=False, with_law=False, with_charts=True, charts_dir=tmp_path / "_c",
    )
    assert out.exists()
    d = Document(str(out))
    # 시나리오4 × (플로우차트·매트릭스·추가세부담막대·실행타임라인) = 16 이미지
    assert len(d.inline_shapes) == 16
    heads = [p.text for p in d.paragraphs
             if p.style.name.startswith("Heading") or p.style.name == "Title"]
    assert any("종합세무검토" in h for h in heads)
    assert any("회사 개요" in h for h in heads)
    assert any("분개장" in h for h in heads)
    assert any("증빙" in h for h in heads)
    # [사실관계] 라벨이 본문에 등장
    assert any("[사실관계]" in p.text for p in d.paragraphs)
