"""tests/test_cpa_report.py — 회계사 친화 자연어 보고서 렌더러 검증(src.cpa_report).

핵심 보장:
  1. 근거가 **클릭 가능한 하이퍼링크**(국가법령정보센터)로 들어간다.
  2. **Noto Sans KR(Google 폰트)** 가 문서 전체에 적용된다.
  3. 구조적 개발어(열거형/코드/내부식별자: FISCAL_YEAR·ANSWERED·ProvisionVersion·client_…)는
     자연어로 치환·은닉된다.
  4. **깨끗한 한국어 입력**(회계사 산출물)은 본문에 영어 알파벳이 0 이다(개발어/영어 금지).

모두 오프라인(법령 fixture replay) — 실키/네트워크/LLM 불필요.
"""
from __future__ import annotations

import re
from dataclasses import replace
from datetime import date
from pathlib import Path

from contract.cluster_f_qa import SynthesisOpinion
from docx import Document
from docx.oxml.ns import qn

from src.ai.law_data_source import default_law_source
from src.cpa_report import _ko, build_cpa_review_docx
from src.draft import (
    ChannelResult, DraftPackageData, InputMaterial, IssueMemo, LawTraceEntry,
    OpportunityItem, ReasoningStep, ReasoningTrace, RiskItem, StrategyOption,
)
from src.draft_demo import build_demo_draft_package
from src.law_anchor import build_law_source_answer
from src.legal_research import ResearchLiteAnswer
from src.review_items import generate_review_items


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def _visible_text(path: Path) -> str:
    """문단 + 표 셀 + 하이퍼링크 앵커 텍스트까지 모두 모은 가시 텍스트."""
    doc = Document(str(path))

    def ptext(par) -> str:
        s = par.text
        for hl in par._p.findall(qn("w:hyperlink")):
            for t in hl.findall(".//" + qn("w:t")):
                s += t.text or ""
        return s

    chunks = [ptext(p) for p in doc.paragraphs]
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                chunks += [ptext(p) for p in c.paragraphs]
    return "\n".join(chunks)


def _hyperlink_targets(path: Path) -> list[str]:
    doc = Document(str(path))
    return [r.target_ref for r in doc.part.rels.values() if "hyperlink" in r.reltype]


# --------------------------------------------------------------------------- #
# 0) _ko 정화 — 한글에 결합된 개발어/식별자도 영어 0 으로(경계 비의존, codex)
# --------------------------------------------------------------------------- #
def test_ko_strips_korean_glued_jargon():
    """한글에 붙은 개발어(\b 경계가 못 잡는 'CPA검토'·'H1승인'·'client_abc')도 영어 0."""
    samples = [
        "CPA검토 필요", "DOCX보고서로 출력", "ANSWERED상태", "H1승인 후 H5승인까지",
        "내부RAG 연동", "①법령MCP 조회", "client_abc 식별자", "matter_x 사건",
        "예규 vs 심판례", "FISCAL_YEAR 기준", "Reviewer확인", "HALU-008 위반",
        "KICPA 기준 검토", "최종 KICPA가 서명",
    ]
    for s in samples:
        out = _ko(s)
        assert not re.search(r"[A-Za-z]", out), f"{s!r} → {out!r} 에 영어 잔존"


# --------------------------------------------------------------------------- #
# 1) 렌더러 동작 — 데모 패키지(하이퍼링크·폰트·구조적 개발어 치환)
# --------------------------------------------------------------------------- #
def test_cpa_report_hyperlinks_font_and_jargon_mapping(tmp_path):
    data, titles, articles = build_demo_draft_package()
    out = build_cpa_review_docx(data, tmp_path / "cpa.docx", titles=titles, articles=articles)
    body = _visible_text(out)

    # (2) 폰트 — 문서 전체 Noto Sans KR
    doc = Document(str(out))
    assert doc.styles["Normal"].font.name == "Noto Sans KR"

    # (1) 근거 하이퍼링크 — 국가법령정보센터로 클릭 이동
    targets = _hyperlink_targets(out)
    assert targets, "근거 하이퍼링크가 하나도 없음"
    assert all("law.go.kr" in t for t in targets)

    # (3) 구조적 개발어/식별자 치환·은닉 — 회계사 보고서에 노출 금지
    for banned in ("FISCAL_YEAR", "ANSWERED", "SILENT", "ProvisionVersion",
                   "client_", "matter_", "Executive Summary", "ReasoningTrace",
                   "law-tracing", "HITL"):
        assert banned not in body, f"개발어 '{banned}' 가 보고서에 노출됨"

    # 자연어 제목(한국어)
    assert "1. 핵심 요약" in body
    assert "9. 관련 법령" in body
    assert "처리 방향별 비교" in body


# --------------------------------------------------------------------------- #
# 2) 영어 0 보장 — 깨끗한 한국어 입력은 본문에 알파벳이 없어야 한다
# --------------------------------------------------------------------------- #
def _clean_korean_package():
    """오프라인 법령 replay(소득세법 제22조)로 만든 *깨끗한 한국어* 검토패키지(영어 0 입력)."""
    law = default_law_source()
    as_of, lookup = date(2026, 1, 1), date(2024, 1, 1)
    bundle = build_law_source_answer(
        lookup=replace(law.lookup_provision("소득세법", "제22조", lookup), as_of_date=as_of),
        client_id="client_t", answer_run_id="ar_t", source_answer_id="sa_t_c22",
        source_type_label="법률", matter_id="matter_t")
    c22 = bundle.citation
    cid = c22.citation_id
    research = ResearchLiteAnswer(
        bundle=bundle,
        legal_conclusion="현실적 퇴직을 원인으로 받는 대가는 명칭과 무관하게 원칙적으로 퇴직소득이다.",
        certainty="해석", reasoning="제22조 퇴직소득 범위·임원 한도.",
        limits_deadlines="현실적 퇴직일이 속하는 과세기간 귀속.",
        filing_impact="퇴직소득세 원천징수·정산.",
        risks=["임원 한도 초과분 근로소득 의제 가능성"], spotted_issues=["퇴직소득 구분"],
        needs_review=True, abstained=False, answer_text="(검토항목 입력)",
        conclusion_claim=bundle.claim, conclusion_citation=c22)
    review_items = generate_review_items(
        answer=research, client_id="client_t", matter_id="matter_t",
        high_risk=True, aggressive=True, prefix="ri_t")
    synthesis = SynthesisOpinion(
        synthesis_id="syn_t", answer_run_id="ar_t", client_id="client_t",
        source_answer_ids=["sa_t_c22"], policy_version="synth-v1",
        opinion_text="법령 원문으로 핵심 결론을 확정하고, 종합 의견은 새 근거를 만들지 않습니다.",
        abstained=False, authority_deficit=None, confidence_cap=None)
    trace = ReasoningTrace(
        steps=[
            ReasoningStep(1, "INTAKE", "사실관계를 수집했습니다.", []),
            ReasoningStep(2, "RESEARCH_CH1_LAW", "소득세법 제22조 원문을 회수했습니다.",
                          [c22.source_locator]),
            ReasoningStep(3, "SYNTHESIS", "법령 원문으로 결론을 확정했습니다.", [c22.source_locator]),
        ],
        law_trace=[LawTraceEntry(
            issue="퇴직소득 구분", law_name="소득세법", article="제22조",
            as_of=as_of.strftime("%Y-%m-%d"),
            basis_kind=c22.applicable_basis.basis_kind.value,
            locator=c22.source_locator, quote_excerpt=(c22.quote or "").strip()[:100])])
    opt = lambda k, l: StrategyOption(  # noqa: E731
        option_key=k, label=l, expected_tax_burden="가정에 따라 달라짐",
        tax_risk="검토 필요", defensibility="법령 근거로 방어",
        required_evidence="정관·지급규정", cpa_review_point="한도 계산 확인",
        citation_ids=[cid])
    data = DraftPackageData(
        matter_id="matter_t", client_id="client_t", company_name="가나다 주식회사",
        fiscal_year="2025 사업연도", as_of_date=as_of,
        review_scope="임원 퇴직위로금의 소득구분을 검토합니다.", high_risk=True,
        executive_summary="현실적 퇴직을 원인으로 받는 퇴직위로금은 원칙적으로 퇴직소득에 해당합니다.",
        input_materials=[InputMaterial("정관·지급규정", "수집", "L2")],
        risks=[RiskItem("임원 한도 초과분 근로소득 의제",
                        "임원 퇴직소득 한도를 초과한 부분은 근로소득으로 봅니다.",
                        citation_ids=[cid], severity="높음")],
        opportunities=[OpportunityItem("퇴직소득 분류과세 이점",
                       "한도 이내 금액은 퇴직소득으로 분류과세받습니다.", citation_ids=[cid])],
        strategy_options=[opt("보수", "보수적 처리"), opt("중립", "중립적 처리"),
                          opt("적극", "적극적 처리")],
        issue_memos=[IssueMemo(topic="쟁점. 퇴직소득 구분",
                     analysis="명칭과 무관하게 현실적 퇴직 대가는 퇴직소득입니다.",
                     citation_ids=[cid], escalation="임원 한도 계산 선행")],
        citations=[c22],
        additional_requests=["임원 퇴직급여 지급규정"],
        review_items=review_items,
        conclusion="현실적 퇴직 대가는 퇴직소득으로 처리하되 임원 한도를 확인합니다.",
        recommended_order=["임원 여부 확정", "한도 계산", "소득구분 확정"],
        data_limits=["임원 한도 자료 미확정"],
        synthesis=synthesis, source_objects=list(bundle.source_objects),
        channel_results=[
            ChannelResult(channel="①", source_label="국가법령정보 조회", status="ANSWERED",
                          answered=True, answer_excerpt="소득세법 제22조 원문을 회수했습니다.",
                          citation_locators=["소득세법 제22조"]),
            ChannelResult(channel="②", source_label="내부 실무기준 자료", status="SILENT",
                          answered=False, answer_excerpt="연동된 내부 자료가 없습니다.",
                          citation_locators=[]),
            ChannelResult(channel="③", source_label="공식 기관 공개자료", status="SILENT",
                          answered=False, answer_excerpt="공개자료 조사를 수행하지 않았습니다.",
                          citation_locators=[])],
        reasoning_trace=trace)
    titles = {cid: "퇴직소득"}
    articles = {cid: "제22조"}
    return data, titles, articles


def test_cpa_report_clean_korean_input_has_zero_english(tmp_path):
    """회계사 산출물(깨끗한 한국어 입력)은 본문 가시 텍스트에 영어 알파벳이 0 이어야 한다.

    하이퍼링크 URL(law.go.kr)은 관계(rel) 대상이라 본문 가시 텍스트에 포함되지 않는다 —
    앵커는 한국어 법령명으로 표시된다."""
    data, titles, articles = _clean_korean_package()
    out = build_cpa_review_docx(data, tmp_path / "clean.docx", titles=titles, articles=articles)
    body = _visible_text(out)
    english = sorted(set(re.findall(r"[A-Za-z]+", body)))
    assert english == [], f"회계사 보고서 본문에 영어 노출: {english}"
    # 그래도 근거 하이퍼링크는 살아 있어야 한다
    assert any("law.go.kr" in t for t in _hyperlink_targets(out))
