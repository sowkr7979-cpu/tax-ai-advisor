"""src/scenario_report.py — '경우의 수' 절세전략·세무리스크 의사결정 보고서 DOCX 렌더러.

``scenario_planner`` 의 시나리오(거래/사건별 분기 추론)를 회계사 친화 DOCX 로 렌더한다.
각 시나리오마다 (1) [사실관계]·추론요지, (2) **경우의 수 플로우차트(그림)**, (3) 경우별 결론
매트릭스(그림+표), (4) 권고안, (5) **사후관리·필요조치·실행계획**, (6) 회계처리 분개 예시,
(7) 내부 RAG DB 근거, (8) 근거 법령 하이퍼링크를 싣고, 마지막에 **분개장·원장 dummy** 부록을
붙인다.

두 가지 산출:
* ``build_scenario_docx`` — 시나리오 묶음만(일반 데모).
* ``build_company_case_docx`` — DART 실재무 그라운딩 + **종합세무검토** + 증빙 dummy 포함(회사 케이스).

표현 헬퍼(_ko·하이퍼링크·폰트·음영·문서스타일)는 ``src.cpa_report`` 의 것을 재사용한다. 근거는
시나리오 모델이 강제(basis 필수)하며, 본 산출물은 **인공지능 추론 초안**으로 회계사 검토(HITL)를
전제로 한다(영어 0 · Noto Sans KR).
"""
from __future__ import annotations

import urllib.parse
from pathlib import Path
from typing import Optional

from src.cpa_report import (
    _add_hyperlink,
    _apply_document_style,
    _cell,
    _ko,
    _set_run_font,
    _shade,
)

_GREY = (0x5F, 0x63, 0x68)
_HEADING_RGB = (0x1A, 0x73, 0xE8)
_HEADER_FILL = "E8F0FE"
_RISK_FILL = {"안전": "E6F4EA", "주의": "FEF7E0", "위험": "FCE8E6"}
_RISK_RGB = {"안전": (0x18, 0x80, 0x38), "주의": (0xE8, 0x71, 0x0A), "위험": (0xD9, 0x30, 0x25)}


def _law_search_url(label: str) -> str:
    """근거 법령 라벨 → 국가법령정보센터 검색 링크(특정 조문 permalink 날조 대신 검색)."""
    law = label.split("(")[0].strip()
    q = urllib.parse.quote(law)
    return f"https://www.law.go.kr/LSW/lsSc.do?menuId=1&subMenuId=15&query={q}"


def _won(v: float) -> str:
    return f"{v:,.0f}"


def _rag_evidence(query: str, *, top_k: int = 2) -> list[tuple[str, str]]:
    """내부 RAG DB 인덱스에서 관련 근거 passage 회수(있으면). 미가용 시 빈 목록(degrade·날조0)."""
    try:
        import src.rag_db_index as ragdb

        index = ragdb.load_index()
        res = ragdb.query_rag_db(index, query, top_k=top_k)
        return [(p.source, p.text.strip().replace("\n", " ")[:180]) for p in res.passages]
    except Exception as exc:  # noqa: BLE001 - 인덱스 미가용은 정상 degrade(날조 금지)
        # fail-closed([]는 유지)하되 실패를 침묵하지 않고 운영자에게 알린다(codex 적발).
        import sys

        print(f"  ⚠ 내부 RAG DB 회수 불가({type(exc).__name__}) — 내부자료 근거 생략: {query[:40]}",
              file=sys.stderr)
        return []


# --------------------------------------------------------------------------- #
# 문서 저수준 헬퍼(doc 인자) — 두 빌더가 공유
# --------------------------------------------------------------------------- #
def _H(doc, txt: str, level: int = 1):
    doc.add_heading(_ko(txt), level=level)


def _body(doc, txt: str, *, size: float = 10.5, rgb=None, bold=False):
    p = doc.add_paragraph()
    _set_run_font(p.add_run(_ko(txt)), size=size, rgb=rgb, bold=bold)
    return p


def _bullets(doc, items, *, size: float = 10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        _set_run_font(p.add_run(_ko(it)), size=size)


def _add_img(doc, path, width=6.6):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    if path and Path(path).exists():
        doc.add_picture(str(path), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _label_para(doc, label: str, text: str):
    p = doc.add_paragraph()
    _set_run_font(p.add_run(label + " "), size=10.5, bold=True, rgb=_HEADING_RGB)
    _set_run_font(p.add_run(_ko(text)), size=10.5)
    return p


# --------------------------------------------------------------------------- #
# 시나리오 1건 렌더(두 빌더 공유) — [사실관계]로 라벨
# --------------------------------------------------------------------------- #
def _render_scenario(doc, n: int, s, flow_path, mtx_path, *, with_rag: bool) -> None:
    _H(doc, f"{n}. {s.title}")
    _label_para(doc, "[사실관계]", s.facts or s.trigger)
    _label_para(doc, "[추론 요지]", s.summary)

    # (1) 플로우차트
    _H(doc, f"{n}-1. 경우의 수 의사결정 플로우차트", level=2)
    _add_img(doc, flow_path, 6.9)

    # (2) 경우의 수별 결론 — 매트릭스 그림 + 표
    _H(doc, f"{n}-2. 경우의 수별 결론 (세무리스크 · 절세전략)", level=2)
    _add_img(doc, mtx_path, 6.9)
    tbl = doc.add_table(rows=1, cols=5); tbl.style = "Table Grid"
    for i, h in enumerate(["경우의 수", "위험", "세무리스크", "절세전략·리스크관리", "근거"]):
        _cell(tbl.rows[0].cells[i], h, header=True, size=9.5)
    for lf in s.leaves:
        r = tbl.add_row().cells
        fill = _RISK_FILL.get(lf.risk_level)
        _cell(r[0], f"{lf.case_label}\n{lf.path}", size=9)
        if fill:
            for c in r:
                _shade(c, fill)
        _cell(r[1], lf.risk_level, size=9, bold=True)
        _set_run_font(r[1].paragraphs[0].runs[0], size=9, bold=True,
                      rgb=_RISK_RGB.get(lf.risk_level))
        _cell(r[2], _ko(lf.tax_risk), size=8.8)
        _cell(r[3], _ko(lf.management), size=8.8)
        r[4].text = ""
        pp = r[4].paragraphs[0]
        for j, b in enumerate(lf.basis):
            if j:
                _set_run_font(pp.add_run(" "), size=8.3)
            _add_hyperlink(pp, _law_search_url(b), b.split("(")[0].strip(), size=8.3)

    # (3) 권고안
    rec = next((lf for lf in s.leaves if lf.case_label == s.recommended_case), None)
    _H(doc, f"{n}-3. 권고안", level=2)
    if rec:
        pr = doc.add_paragraph()
        _set_run_font(pr.add_run(f"▶ 권고: {rec.case_label} ({rec.path}) — 위험등급 ‘{rec.risk_level}’"),
                      size=10.5, bold=True, rgb=_RISK_RGB.get(rec.risk_level, _HEADING_RGB))
        _body(doc, f"{rec.management} 이 경로를 따르면 {s.safe_title}.")

    # (4) 사후관리 · 필요조치 · 실행계획
    _H(doc, f"{n}-4. 사후관리 · 필요조치 · 실행계획", level=2)
    _body(doc, "[사후관리] 거래 종결 후 지속 점검 사항", size=10, bold=True, rgb=_HEADING_RGB)
    _bullets(doc, s.aftercare, size=10)
    _body(doc, "[필요조치] 지금 즉시 해야 할 일", size=10, bold=True, rgb=_HEADING_RGB)
    _bullets(doc, s.actions, size=10)
    _body(doc, "[실행계획] 시점별 실행 로드맵 (미래 plan)", size=10, bold=True, rgb=_HEADING_RGB)
    pt = doc.add_table(rows=1, cols=3); pt.style = "Table Grid"
    for i, h in enumerate(["시점", "행위", "세무 효과·목적"]):
        _cell(pt.rows[0].cells[i], h, header=True, size=9.5)
    for st in s.plan:
        r = pt.add_row().cells
        _cell(r[0], st.when, size=9.3, bold=True); _cell(r[1], _ko(st.what), size=9.3)
        _cell(r[2], _ko(st.effect), size=9.3)

    # (5) 회계처리 분개 예시
    if s.journal:
        _H(doc, f"{n}-5. 회계처리 분개 예시 (권고 경로)", level=2)
        _body(doc, s.journal.title, size=10, bold=True)
        _body(doc, s.journal.note, size=9.3, rgb=_GREY)
        jt = doc.add_table(rows=1, cols=3); jt.style = "Table Grid"
        for i, h in enumerate(["계정과목", "차변(원)", "대변(원)"]):
            _cell(jt.rows[0].cells[i], h, header=True, size=9.5)
        for ln in s.journal.lines:
            r = jt.add_row().cells
            _cell(r[0], _ko(ln.account), size=9.3)
            _cell(r[1], _won(ln.debit) if ln.debit else "", size=9.3)
            _cell(r[2], _won(ln.credit) if ln.credit else "", size=9.3)

    # (6) 내부 자료(RAG DB) 근거
    if with_rag:
        ev = _rag_evidence(s.rag_query or f"{s.title} 세무 리스크 절세", top_k=2)
        if ev:
            _H(doc, f"{n}-6. 내부 자료 근거 (실무 가이드 회수)", level=2)
            _body(doc, "내부 실무자료 임베딩 색인에서 회수한 관련 근거입니다(출처·발췌).",
                  size=9.3, rgb=_GREY)
            for src_loc, excerpt in ev:
                bp = doc.add_paragraph(style="List Bullet")
                _set_run_font(bp.add_run(f"[{src_loc}] "), size=9, bold=True, rgb=_GREY)
                _set_run_font(bp.add_run(_ko(excerpt) + " …"), size=9)

    # (7) 근거 법령
    _H(doc, f"{n}-7. 근거 법령 (클릭하면 검색)", level=2)
    for c in s.citations:
        bp = doc.add_paragraph(style="List Bullet")
        _add_hyperlink(bp, _law_search_url(c), c, size=10)


# --------------------------------------------------------------------------- #
def _new_doc():
    from datetime import datetime

    from docx import Document

    doc = Document()
    doc.core_properties.created = datetime(2024, 1, 1)
    doc.core_properties.modified = datetime(2024, 1, 1)
    _apply_document_style(doc)
    return doc


def _render_charts(scenarios, charts_dir, with_charts):
    flow_paths: dict[str, Path] = {}
    mtx_paths: dict[str, Path] = {}
    if with_charts:
        from src import scenario_flowchart as sfc

        for s in scenarios:
            flow_paths[s.key] = sfc.save_scenario_flowchart(s, charts_dir / f"flow_{s.key}.png")
            mtx_paths[s.key] = sfc.save_case_matrix(s, charts_dir / f"matrix_{s.key}.png")
    return flow_paths, mtx_paths


def build_scenario_docx(
    scenarios,
    out_path: str | Path,
    *,
    company: str = "가나다정밀(주)",
    as_of: str = "2026년 6월 30일",
    charts_dir: Optional[str | Path] = None,
    with_charts: bool = True,
    with_rag: bool = True,
) -> Path:
    """경우의 수 의사결정 보고서 DOCX(플로우차트·매트릭스·사후관리/조치/계획·분개 dummy)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for s in scenarios:
        s.validate()
    out_path = Path(out_path)
    charts_dir = Path(charts_dir) if charts_dir else out_path.parent / "_charts_scn"
    flow_paths, mtx_paths = _render_charts(scenarios, charts_dir, with_charts)

    doc = _new_doc()
    t = doc.add_heading("경우의 수 기반 절세전략·세무리스크 의사결정 보고서", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(sub.add_run(f"{company}　·　검토기준일 {as_of}　·　내부 검토용(인공지능 추론 초안)"),
                  size=10, rgb=_GREY)
    _render_method_note(doc)

    for n, s in enumerate(scenarios, start=1):
        _render_scenario(doc, n, s, flow_paths.get(s.key), mtx_paths.get(s.key), with_rag=with_rag)

    _append_books(doc, scenarios)
    _render_closing(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def build_company_case_docx(
    scenarios,
    out_path: str | Path,
    *,
    company: str,
    as_of: str,
    overview_rows: list[tuple[str, str]],
    overview_note: str,
    tax_review_rows: list[tuple[str, str, str, str]],
    evidence: list[str],
    data_source_note: str,
    charts_dir: Optional[str | Path] = None,
    with_charts: bool = True,
    with_rag: bool = True,
) -> Path:
    """DART 실재무 그라운딩 회사 케이스 보고서 — 회사개요 + **종합세무검토** + 시나리오 + 분개장/원장/증빙 dummy."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for s in scenarios:
        s.validate()
    out_path = Path(out_path)
    charts_dir = Path(charts_dir) if charts_dir else out_path.parent / "_charts_case"
    flow_paths, mtx_paths = _render_charts(scenarios, charts_dir, with_charts)

    doc = _new_doc()
    t = doc.add_heading(f"{company} 종합세무검토 및 경우의 수 절세전략·세무리스크 의사결정 보고서", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(sub.add_run(f"{company}　·　검토기준일 {as_of}　·　내부 검토용(인공지능 추론 초안)"),
                  size=10, rgb=_GREY)
    note = doc.add_paragraph(); note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(note.add_run(_ko(data_source_note)), size=9.3, rgb=(0xB0, 0x60, 0x00))

    # 0. 방법론
    _render_method_note(doc)

    # 1. 회사 개요 및 사실관계 기초(DART 실재무)
    _H(doc, "1. 회사 개요 및 사실관계 기초 (DART 전자공시)")
    ot = doc.add_table(rows=1, cols=2); ot.style = "Table Grid"
    _cell(ot.rows[0].cells[0], "항목", header=True, size=9.7)
    _cell(ot.rows[0].cells[1], "값(DART 공시 기준)", header=True, size=9.7)
    for k, v in overview_rows:
        r = ot.add_row().cells
        _cell(r[0], k, size=9.5, bold=True); _cell(r[1], v, size=9.5)
    _body(doc, overview_note, size=10)

    # 2. 종합세무검토(재무제표 → 잠재 세무쟁점)
    _H(doc, "2. 종합세무검토 (재무제표에서 도출한 주요 세무 쟁점)")
    _body(doc, "DART 공시 재무수치를 단서로, 세무조사·신고 시 쟁점이 될 수 있는 항목과 검토 "
               "방향을 정리했습니다. 상세 분기 추론은 3장 이하 시나리오에서 다룹니다.", size=9.7, rgb=_GREY)
    rt = doc.add_table(rows=1, cols=4); rt.style = "Table Grid"
    for i, h in enumerate(["계정·근거 수치", "잠재 세무쟁점", "검토 방향", "근거"]):
        _cell(rt.rows[0].cells[i], h, header=True, size=9.5)
    for acct, issue, direction, basis in tax_review_rows:
        r = rt.add_row().cells
        _cell(r[0], _ko(acct), size=8.8, bold=True); _cell(r[1], _ko(issue), size=8.8)
        _cell(r[2], _ko(direction), size=8.8)
        r[3].text = ""
        _add_hyperlink(r[3].paragraphs[0], _law_search_url(basis), basis.split("(")[0].strip(), size=8.5)

    # 3..N. 시나리오별 의사결정
    for n, s in enumerate(scenarios, start=3):
        _render_scenario(doc, n, s, flow_paths.get(s.key), mtx_paths.get(s.key), with_rag=with_rag)

    # 부록 분개장·원장 + 증빙
    _append_books(doc, scenarios)
    _append_evidence(doc, evidence)
    _render_closing(doc)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# --------------------------------------------------------------------------- #
def _render_method_note(doc) -> None:
    _H(doc, "0. 이 보고서를 읽는 법 (방법론)")
    _body(doc, "실무 세무 판단은 ‘이 거래는 안전한가/위험한가’라는 단답이 아니라, 몇 개의 핵심 "
               "조건을 어떻게 충족하느냐에 따라 결과가 갈리는 ‘경우의 수’ 문제입니다. 본 보고서는 각 "
               "거래를 순차 분기(판단 게이트)로 분해하고, 분기마다 ① 빗나갔을 때의 세무리스크(과세·"
               "손금부인·추징)와 ② 통과시키기 위한 절세전략·리스크관리(요건·증빙·시점)를 함께 제시합니다.")
    _body(doc, "플로우차트 색상 규칙 — 주황 마름모: 판단(분기) / 빨강: 세무리스크(빗나간 경우) / "
               "초록: 정상 종결(요건 충족). 세로 화살표를 따라 모든 게이트를 통과하면 가장 안전한 "
               "경우(권고안)에 도달합니다.", rgb=_GREY, size=9.8)
    _body(doc, "⚠ 본 산출물은 공개 법령·실무 자료에 근거한 인공지능 추론 초안이며, 회사의 재무수치 외 "
               "분개·원장·증빙·거래 사실관계는 검토 예시로 가상 생성(dummy)한 것입니다. 실제 적용 전 "
               "반드시 담당 회계사의 사실관계 확인·검토가 필요합니다.", rgb=(0xB0, 0x60, 0x00), size=9.8)


def _render_closing(doc) -> None:
    _H(doc, "종합 유의사항")
    _bullets(doc, [
        "본 보고서의 분기·결론은 일반적 사실관계 가정에 따른 인공지능 추론이며, 개별 사안의 구체적 "
        "사실관계·최신 예규·판례에 따라 결론이 달라질 수 있습니다.",
        "회사명·재무수치는 DART 전자공시 기준이나, 분개·원장·증빙·거래 사실관계는 추론을 위한 "
        "가상 예시(dummy)이며 실제 회계처리·신고 전 담당 회계사의 검토와 서명이 필요합니다.",
        "근거 법령 링크는 국가법령정보센터 검색으로 연결되며, 적용 시점(시행일) 확인이 필요합니다.",
    ])


def _append_books(doc, scenarios) -> None:
    """전체 시나리오의 분개를 모아 분개장(journal)·원장(ledger) dummy 부록 작성."""
    doc.add_page_break()
    _H(doc, "부록 A. 분개장 (가상 dummy)")
    _body(doc, "권고 경로의 회계처리 분개를 한 곳에 모은 가상의 분개장입니다(예시 일자·금액).",
          size=9.5, rgb=_GREY)
    dates = ["2026-03-10", "2026-04-15", "2026-05-20", "2026-06-12", "2026-07-08"]
    jt = doc.add_table(rows=1, cols=5); jt.style = "Table Grid"
    for i, h in enumerate(["일자", "적요", "계정과목", "차변(원)", "대변(원)"]):
        _cell(jt.rows[0].cells[i], h, header=True, size=9.3)
    ledger: dict[str, list[float]] = {}
    for idx, s in enumerate(scenarios):
        if not s.journal:
            continue
        date = dates[idx % len(dates)]
        for k, ln in enumerate(s.journal.lines):
            r = jt.add_row().cells
            _cell(r[0], date if k == 0 else "", size=8.8)
            _cell(r[1], _ko(s.title) if k == 0 else "", size=8.8)
            _cell(r[2], _ko(ln.account), size=8.8)
            _cell(r[3], _won(ln.debit) if ln.debit else "", size=8.8)
            _cell(r[4], _won(ln.credit) if ln.credit else "", size=8.8)
            acc = _ko(ln.account)
            agg = ledger.setdefault(acc, [0.0, 0.0])
            agg[0] += ln.debit; agg[1] += ln.credit

    _H(doc, "부록 B. 총계정원장 (가상 dummy)")
    _body(doc, "분개장을 계정과목별로 집계한 원장입니다(차변·대변 합계와 잔액).", size=9.5, rgb=_GREY)
    lt = doc.add_table(rows=1, cols=4); lt.style = "Table Grid"
    for i, h in enumerate(["계정과목", "차변 합계(원)", "대변 합계(원)", "잔액(원)"]):
        _cell(lt.rows[0].cells[i], h, header=True, size=9.3)
    tot_d = tot_c = 0.0
    for acc, (d, c) in ledger.items():
        r = lt.add_row().cells
        _cell(r[0], acc, size=8.8)
        _cell(r[1], _won(d) if d else "", size=8.8)
        _cell(r[2], _won(c) if c else "", size=8.8)
        _cell(r[3], f"{d - c:,.0f}", size=8.8)
        tot_d += d; tot_c += c
    r = lt.add_row().cells
    _cell(r[0], "합계", header=True, size=9); _cell(r[1], _won(tot_d), header=True, size=9)
    _cell(r[2], _won(tot_c), header=True, size=9)
    _cell(r[3], f"{tot_d - tot_c:,.0f}", header=True, size=9)
    _body(doc, f"차변 합계 {_won(tot_d)}원 = 대변 합계 {_won(tot_c)}원 "
               f"(대차 일치 {'✓' if abs(tot_d - tot_c) < 1 else '✗'}).", size=9.3, rgb=_GREY)


def _append_evidence(doc, evidence) -> None:
    if not evidence:
        return
    _H(doc, "부록 C. 증빙 목록 (가상 dummy)")
    _body(doc, "각 시나리오의 권고 경로를 입증하기 위해 구비·보관해야 할 증빙의 가상 목록입니다.",
          size=9.5, rgb=_GREY)
    et = doc.add_table(rows=1, cols=3); et.style = "Table Grid"
    for i, h in enumerate(["번호", "증빙명", "용도·보관"]):
        _cell(et.rows[0].cells[i], h, header=True, size=9.3)
    for i, ev in enumerate(evidence, start=1):
        name, _, use = ev.partition("|")
        r = et.add_row().cells
        _cell(r[0], f"증{i:02d}", size=8.8, bold=True)
        _cell(r[1], _ko(name.strip()), size=8.8)
        _cell(r[2], _ko(use.strip()) if use else "—", size=8.8)
