"""src/strategy_plan_report.py — 회사 다(多)전략 절세 종합검토 회계사용 DOCX 렌더러.

``StrategyPlanData`` 를 (1)재무제표(DART 원본·원 단위·쟁점 강조), (2)절세전략 종합 스크리닝표
(적용/조건부/미해당), (3)적용 전략 카테고리별 상세(작동원리·요건·효과·리스크·검토포인트·근거
링크), (4)그림·도표(스크리닝 맵·기대효과 등급·실행 로드맵), (5)근거 법령 하이퍼링크 로 렌더한다.
회계사 자연어(영어·개발어 0)·Noto Sans KR. 표현 헬퍼는 ``src.cpa_report``·``src.tax_plan_report``
의 것을 재사용한다. 인용은 등록 소스객체로 전수 검증(날조 차단·fail-closed)한다.
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from src.cpa_report import (
    _FONT, _HEADING_RGB, _add_hyperlink, _apply_document_style, _ko, _qn,
    _set_run_font, _shade,
)
from src.draft import CitationView
from src.source_registry import CitationVerificationError, SourceRegistry
from src.strategy_plan import StrategyPlanData
from src.tax_plan_report import (
    _HEADER_FILL, _ISSUE_FILL, _SUBHEAD_FILL, _font_cell, _issue_cell,
)

_GREY = (0x5F, 0x63, 0x68)
_APPLY_FILL = "E6F4EA"      # 적용(연한 그린)
_COND_FILL = "FEF7E0"       # 조건부(연한 노랑)
_NA_FILL = "F1F3F4"         # 미해당(연한 회색)

_SECTIONS = [
    "1. 핵심 요약 및 목표",
    "2. 법령·RAG·웹 3소스 교차검증 및 종합의견",
    "3. 회사 개요",
    "4. 재무상태표 (전자공시 원본 · 쟁점 강조)",
    "5. 손익계산서 (전자공시 원본 · 쟁점 강조)",
    "6. 계정별원장 (세무 관련 주요 계정 · 시연용 예시)",
    "7. 분개장 (세무 쟁점 거래 · 시연용 예시)",
    "8. 절세전략 종합 스크리닝 (전체 전략 × 적용 여부)",
    "9. 적용 권고 전략 상세 (필요 조치 · 자료 · 사후관리 포함)",
    "10. 실행 우선순위 및 로드맵",
    "11. 관련 법령 및 근거 (클릭하면 원문으로 이동)",
    "12. 회계사 검토 필요사항",
    "13. 자료·가정의 한계 (정직 고지)",
    "14. 결론",
]


class StrategyPlanValidationError(RuntimeError):
    """인용 날조/미해소 — 생성 실패(fail-closed)."""


def _status_label(card) -> tuple[str, str]:
    """(라벨, 음영색) — 적용/조건부/미해당."""
    if not card.applies:
        return "미해당", _NA_FILL
    if "조건부" in card.applicability_note:
        return "조건부", _COND_FILL
    return "적용", _APPLY_FILL


def _validate(plan: StrategyPlanData) -> None:
    cidx = plan.citation_index
    registry = SourceRegistry() if plan.source_objects else None
    if registry is not None:
        registry.register_all(plan.source_objects)

    if not plan.citations:
        raise StrategyPlanValidationError("근거 법령(Citation) 없음")

    # (A) 렌더되는 모든 인용을 등록 소스객체로 검증.
    for cit in plan.citations:
        if registry is None:
            raise StrategyPlanValidationError(
                f"인용 미검증(날조 차단): {cit.citation_id} — source_objects 필요")
        try:
            registry.require_citation(cit)
        except CitationVerificationError as exc:
            raise StrategyPlanValidationError(f"인용 검증 실패: {cit.citation_id} — {exc}") from exc

    # (B) 본문 참조 citation_id 가 전부 검증된 인용으로 해소되는지 확인.
    def check_ids(ids, where):
        for cid in ids:
            if cid is None or cidx.get(cid) is None:
                raise StrategyPlanValidationError(f"인용 미해소: {where} 의 {cid}")

    for ln in (*plan.balance_sheet.lines, *plan.income_statement.lines):
        if ln.is_issue and ln.citation_ids:
            check_ids(ln.citation_ids, f"재무제표 쟁점행 '{ln.label}'")
    for st in plan.timeline:
        check_ids([c for c in st.citation_ids if c is not None], f"타임라인 '{st.action}'")
    for s in plan.strategies:
        for (law_name, art) in s.citation_articles:
            cid = plan.art_to_cid.get((law_name, art))
            if cid is None or cidx.get(cid) is None:
                raise StrategyPlanValidationError(f"인용 미해소: 전략 '{s.title}' 의 {law_name} {art}")
    # 3소스 리서치 채널의 법령 인용도 등록 인용으로 전수 해소(날조 차단).
    research = getattr(plan, "research", None)
    if research is not None:
        for ch in research.channels:
            for (law_name, art) in getattr(ch, "citation_articles", []):
                cid = plan.art_to_cid.get((law_name, art))
                if cid is None or cidx.get(cid) is None:
                    raise StrategyPlanValidationError(
                        f"인용 미해소: 리서치 채널 '{ch.channel}' 의 {law_name} {art}")
    if not any(s.applies for s in plan.strategies):
        raise StrategyPlanValidationError("적용 전략(applies) 1건 이상 필요")


def _cviews_for(plan: StrategyPlanData, titles, articles):
    return {cv.citation_id: cv for cv in
            [CitationView.from_citation(c, title=(titles or {}).get(c.citation_id, ""),
                                        article_label=(articles or {}).get(c.citation_id, ""))
             for c in plan.citations]}


def _strategy_cids(plan: StrategyPlanData, card) -> list[str]:
    return [plan.art_to_cid[t] for t in card.citation_articles if t in plan.art_to_cid]


def _strategy_cids_by_key(plan: StrategyPlanData, key: str) -> list[str]:
    for s in plan.strategies:
        if s.key == key:
            return _strategy_cids(plan, s)
    return []


def _channel_cids(plan: StrategyPlanData, ch) -> list[str]:
    return [plan.art_to_cid[t] for t in ch.citation_articles if t in plan.art_to_cid]


def _render_research(doc, plan: StrategyPlanData, cviews, cite_runs) -> None:
    """§2 — 법령①·내부자료②·공식웹③ 3소스 답변 + 종합의견. ②는 실제 임베딩 인덱스 회수
    근거(가이드 파일·페이지)를 그대로 노출한다(없으면 보류). 법령 인용은 하이퍼링크로 검증."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: F401 (도식 정렬 일관)

    research = getattr(plan, "research", None)
    if research is None:
        p = doc.add_paragraph()
        _set_run_font(p.add_run("(3소스 리서치 정보가 제공되지 않았습니다.)"), size=10, rgb=_GREY)
        return

    qp = doc.add_paragraph()
    _set_run_font(qp.add_run("[검토 질문] "), size=10.5, bold=True, rgb=_HEADING_RGB)
    _set_run_font(qp.add_run(_ko(research.question)), size=10.5)

    for ch in research.channels:
        hp = doc.add_paragraph()
        _set_run_font(hp.add_run(f"▶ {_ko(ch.channel)} "), size=10.5, bold=True, rgb=_HEADING_RGB)
        _set_run_font(hp.add_run(f"［{_ko(ch.authority)}］"), size=9, rgb=_GREY)
        if getattr(ch, "abstained", False):
            _set_run_font(hp.add_run("  — 보류(근거 없음·침묵)"), size=9, rgb=(0xB0, 0x60, 0x00))
        ap = doc.add_paragraph()
        _set_run_font(ap.add_run(_ko(ch.answer)), size=10)

        cids = _channel_cids(plan, ch)
        if cids:
            gp = doc.add_paragraph()
            _set_run_font(gp.add_run("　근거(법령): "), size=9, rgb=_GREY)
            cite_runs(gp, cids, size=9)
        if getattr(ch, "web_url", ""):
            wp = doc.add_paragraph()
            _set_run_font(wp.add_run("　공식 출처: "), size=9, rgb=_GREY)
            _add_hyperlink(wp, ch.web_url, "국가법령정보센터 원문", size=9)
        for ev in getattr(ch, "rag_evidence", []):
            ep = doc.add_paragraph(style="List Bullet")
            _set_run_font(ep.add_run(f"［{_ko(ev.source)}］ "), size=8.6, bold=True, rgb=_HEADING_RGB)
            excerpt = ev.text.strip().replace("\n", " ")
            excerpt = excerpt[:240] + "…" if len(excerpt) > 240 else excerpt
            _set_run_font(ep.add_run(_ko(excerpt)), size=8.6)
        note = getattr(ch, "note", "")
        if note and not getattr(ch, "rag_evidence", []):
            np_ = doc.add_paragraph()
            _set_run_font(np_.add_run("　▸ " + _ko(note)), size=8.5, rgb=_GREY)

    op = doc.add_paragraph()
    _set_run_font(op.add_run("【종합의견】 "), size=10.5, bold=True, rgb=(0x18, 0x80, 0x38))
    _set_run_font(op.add_run(_ko(research.opinion)), size=10.5)
    if getattr(research, "method_note", ""):
        mp = doc.add_paragraph()
        _set_run_font(mp.add_run("　▸ " + _ko(research.method_note)), size=8.7, rgb=(0xB0, 0x60, 0x00))


def build_strategy_plan_docx(
    plan: StrategyPlanData,
    out_path: str | Path,
    *,
    titles: Optional[dict] = None,
    articles: Optional[dict] = None,
    charts_dir: Optional[str | Path] = None,
    with_charts: bool = True,
) -> Path:
    """다전략 종합검토 DOCX(스크리닝·쟁점 강조·도표·하이퍼링크·Noto Sans KR·영어 0)."""
    from datetime import datetime

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    _validate(plan)
    cviews = _cviews_for(plan, titles, articles)

    out_path = Path(out_path)
    charts_dir = Path(charts_dir) if charts_dir else out_path.parent / "_charts"
    chart_paths: dict[str, Path] = {}
    if with_charts:
        from src import tax_charts as tc
        chart_paths["screen"] = tc.save_strategy_screening(
            plan.strategies, charts_dir / "screen.png")
        chart_paths["effect"] = tc.save_strategy_effect_tier(
            plan.strategies, charts_dir / "effect.png")
        chart_paths["timeline"] = tc.save_timeline(plan.timeline, charts_dir / "timeline.png")

    doc = Document()
    doc.core_properties.created = datetime(2024, 1, 1)
    doc.core_properties.modified = datetime(2024, 1, 1)
    _apply_document_style(doc)

    # 표지
    t = doc.add_heading(_ko(plan.title), level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(sub.add_run(f"검토기준일 {plan.as_of:%Y년 %m월 %d일}　·　{plan.company_name}　·　내부 검토용"),
                  size=10, rgb=_GREY)
    note = doc.add_paragraph(); note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(note.add_run(_ko(plan.source_note)), size=9, rgb=(0xB0, 0x60, 0x00))

    def H(i): doc.add_heading(_ko(_SECTIONS[i]), level=1)
    def body(txt, size=10.5):
        p = doc.add_paragraph(); _set_run_font(p.add_run(_ko(txt)), size=size); return p

    def add_chart(key, width=6.6):
        if with_charts and key in chart_paths:
            doc.add_picture(str(chart_paths[key]), width=Inches(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def cite_runs(paragraph, cids, *, size=9):
        first = True
        for cid in cids:
            cv = cviews.get(cid)
            if not cv:
                continue
            if not first:
                _set_run_font(paragraph.add_run(", "), size=size)
            _add_hyperlink(paragraph, cv.href, cv.label, size=size)
            first = False

    # 1. 핵심 요약 및 목표
    H(0); body(plan.executive_summary)
    p = doc.add_paragraph(); _set_run_font(p.add_run("[검토 목표] "), size=10.5, bold=True,
                                           rgb=_HEADING_RGB)
    _set_run_font(p.add_run(_ko(plan.objective)), size=10.5)

    # 2. 법령·RAG·웹 3소스 교차검증 및 종합의견
    H(1)
    _render_research(doc, plan, cviews, cite_runs)

    # 3. 회사 개요
    H(2)
    for e in plan.overview:
        bp = doc.add_paragraph(style="List Bullet"); _set_run_font(bp.add_run(_ko(e)), size=10.5)

    # 4·5. 재무제표 (DART 원본, 원 단위, 쟁점 강조)
    for idx, fs in ((3, plan.balance_sheet), (4, plan.income_statement)):
        H(idx)
        cap = doc.add_paragraph()
        _set_run_font(cap.add_run(_ko(f"{fs.period}　({fs.unit})")), size=9.5, rgb=_GREY)
        tbl = doc.add_table(rows=1, cols=3); tbl.style = "Table Grid"
        for i, htxt in enumerate(["계정", "금액(원)", "쟁점·적용 전략"]):
            _font_cell(tbl.rows[0].cells[i], htxt, bold=True, rgb=_HEADING_RGB,
                       fill=_HEADER_FILL, size=10)
        for ln in fs.lines:
            r = tbl.add_row().cells
            if ln.amount is None:  # 구분행
                _font_cell(r[0], ln.label, bold=True, fill=_SUBHEAD_FILL, size=10)
                _font_cell(r[1], "", fill=_SUBHEAD_FILL); _font_cell(r[2], "", fill=_SUBHEAD_FILL)
                continue
            name = ("　" * ln.indent) + ln.label
            fill = _ISSUE_FILL if ln.is_issue else None
            _font_cell(r[0], name, bold=ln.bold, fill=fill, size=9.5)
            _font_cell(r[1], f"{ln.amount:,.0f}", bold=ln.bold, fill=fill, size=9.5)
            if ln.is_issue:
                _issue_cell(r[2], ln.issue_label, cviews, ln.citation_ids)
            else:
                _font_cell(r[2], "", fill=fill)

    # 6. 계정별원장 (세무 관련 주요 계정 · 시연용)
    H(5)
    body("아래 계정별원장은 재무제표의 세무 쟁점을 거래 단위로 보여 주기 위한 시연용 예시입니다"
         "(실제 회사의 원장이 아님).", 9.5)
    for lg in plan.ledgers:
        hp = doc.add_paragraph()
        _set_run_font(hp.add_run(f"［{_ko(lg.account)}］ "), size=10, bold=True, rgb=_HEADING_RGB)
        lcids = _strategy_cids_by_key(plan, lg.strategy_key)
        if lcids:
            _set_run_font(hp.add_run("근거: "), size=8.7, rgb=_GREY)
            cite_runs(hp, lcids, size=8.7)
        lt = doc.add_table(rows=1, cols=5); lt.style = "Table Grid"
        for i, htxt in enumerate(["일자", "적요", "차변", "대변", "잔액"]):
            _font_cell(lt.rows[0].cells[i], htxt, bold=True, fill=_HEADER_FILL, size=9)
        r0 = lt.add_row().cells
        _font_cell(r0[0], "기초", size=8.7); _font_cell(r0[1], "전기이월", size=8.7)
        _font_cell(r0[2], "", size=8.7); _font_cell(r0[3], "", size=8.7)
        _font_cell(r0[4], f"{lg.opening:,.0f}", size=8.7)
        for ll in lg.lines:
            r = lt.add_row().cells
            _font_cell(r[0], ll.date, size=8.7); _font_cell(r[1], ll.summary, size=8.7)
            _font_cell(r[2], f"{ll.debit:,.0f}" if ll.debit else "", size=8.7)
            _font_cell(r[3], f"{ll.credit:,.0f}" if ll.credit else "", size=8.7)
            _font_cell(r[4], f"{ll.balance:,.0f}", size=8.7)
        rc = lt.add_row().cells
        _font_cell(rc[0], "기말", bold=True, fill=_ISSUE_FILL, size=8.7)
        _font_cell(rc[1], "차기이월", bold=True, fill=_ISSUE_FILL, size=8.7)
        _font_cell(rc[2], "", fill=_ISSUE_FILL); _font_cell(rc[3], "", fill=_ISSUE_FILL)
        _font_cell(rc[4], f"{lg.closing:,.0f}", bold=True, fill=_ISSUE_FILL, size=8.7)
        npar = doc.add_paragraph()
        _set_run_font(npar.add_run("　▸ " + _ko(lg.note)), size=9, rgb=(0xB0, 0x60, 0x00))

    # 7. 분개장 (세무 쟁점 거래 · 시연용)
    H(6)
    body("아래 분개장은 세무 쟁점이 되는 거래를 발췌한 시연용 예시입니다(실제 회사의 분개가 아님).", 9.5)
    jt = doc.add_table(rows=1, cols=6); jt.style = "Table Grid"
    for i, htxt in enumerate(["순번", "일자", "적요", "차변(계정·금액)", "대변(계정·금액)", "세무 쟁점·근거"]):
        _font_cell(jt.rows[0].cells[i], htxt, bold=True, rgb=_HEADING_RGB, fill=_HEADER_FILL, size=9)
    for je in plan.journal:
        r = jt.add_row().cells
        _font_cell(r[0], str(je.seq), size=8.7); _font_cell(r[1], je.date, size=8.7)
        _font_cell(r[2], je.summary, size=8.5)
        _font_cell(r[3], f"{je.debit_acct} {je.debit_amt:,.0f}", fill=_ISSUE_FILL, size=8.3)
        _font_cell(r[4], f"{je.credit_acct} {je.credit_amt:,.0f}", fill=_ISSUE_FILL, size=8.3)
        rc = r[5]; rc.text = ""; pp = rc.paragraphs[0]
        _set_run_font(pp.add_run("★ " + _ko(je.issue) + " "), size=8.3, bold=True, rgb=(0xB0, 0x60, 0x00))
        cite_runs(pp, _strategy_cids_by_key(plan, je.strategy_key), size=8.3)
        _shade(rc, _ISSUE_FILL)

    # 8. 절세전략 종합 스크리닝
    H(7)
    body(f"재무제표에 대해 법인 절세전략 {len(plan.strategies)}종을 스크리닝했습니다. "
         "'미해당'도 제외하지 않고 사유와 함께 표기합니다(검토 투명성).", 9.5)
    add_chart("screen", 6.6)
    st = doc.add_table(rows=1, cols=5); st.style = "Table Grid"
    for i, htxt in enumerate(["전략", "카테고리", "적용 여부", "스크리닝 사유", "근거"]):
        _font_cell(st.rows[0].cells[i], htxt, bold=True, rgb=_HEADING_RGB, fill=_HEADER_FILL, size=9.5)
    for s in plan.strategies:
        label, fill = _status_label(s)
        r = st.add_row().cells
        _font_cell(r[0], s.title, bold=True, fill=fill, size=9)
        _font_cell(r[1], s.category, fill=fill, size=8.7)
        _font_cell(r[2], label, bold=True, fill=fill, size=9,
                   rgb=(0x18, 0x80, 0x38) if label != "미해당" else _GREY)
        _font_cell(r[3], s.applicability_note, fill=fill, size=8.7)
        # 근거 셀
        rc = r[4]; rc.text = ""; pp = rc.paragraphs[0]
        cids = _strategy_cids(plan, s)
        if cids:
            cite_runs(pp, cids, size=8.5)
        elif s.note_only_basis:
            _set_run_font(pp.add_run("원문 확인 필요"), size=8.5, rgb=(0xB0, 0x60, 0x00))
        else:
            _set_run_font(pp.add_run("—"), size=8.5)
        if fill:
            _shade(rc, fill)

    # 9. 적용 권고 전략 상세 (카테고리별 · 필요 조치·자료·사후관리 포함)
    H(8)
    add_chart("effect", 6.6)
    applied = [s for s in plan.strategies if s.applies]
    cats: list[str] = []
    for s in applied:
        if s.category not in cats:
            cats.append(s.category)
    for cat in cats:
        doc.add_heading(_ko(cat), level=2)
        for s in [x for x in applied if x.category == cat]:
            label, _ = _status_label(s)
            hp = doc.add_paragraph()
            _set_run_font(hp.add_run(f"▶ {_ko(s.title)} "), size=10.5, bold=True, rgb=_HEADING_RGB)
            _set_run_font(hp.add_run(f"[{label}]"), size=9,
                          rgb=(0x18, 0x80, 0x38) if label == "적용" else (0xB0, 0x60, 0x00))
            for tag, val in (("작동 원리", s.mechanism), ("적용 요건", s.requirement),
                             ("기대 효과", s.effect), ("유의·리스크", s.risk)):
                bp = doc.add_paragraph(style="List Bullet")
                _set_run_font(bp.add_run(f"{tag}: "), size=9.8, bold=True)
                _set_run_font(bp.add_run(_ko(val)), size=9.8)
            for tag, items in (("필요 조치", s.actions), ("필요 자료", s.required_data)):
                if items:
                    bp = doc.add_paragraph(style="List Bullet")
                    _set_run_font(bp.add_run(f"{tag}: "), size=9.8, bold=True, rgb=_HEADING_RGB)
                    _set_run_font(bp.add_run(_ko(" / ".join(items))), size=9.8)
            if s.post_management:
                bp = doc.add_paragraph(style="List Bullet")
                _set_run_font(bp.add_run("사후관리: "), size=9.8, bold=True, rgb=(0xB0, 0x60, 0x00))
                _set_run_font(bp.add_run(_ko(s.post_management)), size=9.8)
            if s.review_points:
                bp = doc.add_paragraph(style="List Bullet")
                _set_run_font(bp.add_run("검토 포인트: "), size=9.8, bold=True)
                _set_run_font(bp.add_run(_ko(" / ".join(s.review_points))), size=9.8)
            gp = doc.add_paragraph()
            _set_run_font(gp.add_run("　근거: "), size=9, rgb=_GREY)
            cids = _strategy_cids(plan, s)
            if cids:
                cite_runs(gp, cids, size=9)
            elif s.note_only_basis:
                _set_run_font(gp.add_run(_ko(s.note_only_basis)), size=8.7, rgb=(0xB0, 0x60, 0x00))

    # 10. 실행 우선순위 및 로드맵
    H(9)
    add_chart("timeline", 6.8)
    tt = doc.add_table(rows=1, cols=4); tt.style = "Table Grid"
    for i, htxt in enumerate(["시점", "행위", "세무 효과", "근거"]):
        _font_cell(tt.rows[0].cells[i], htxt, bold=True, fill=_HEADER_FILL, size=9.5)
    for s in plan.timeline:
        r = tt.add_row().cells
        _font_cell(r[0], s.date_label, bold=True, size=9.3); _font_cell(r[1], s.action, size=9.3)
        _font_cell(r[2], s.tax_effect, size=9.3)
        cids = [c for c in s.citation_ids if c]
        rc = r[3]; rc.text = ""; pp = rc.paragraphs[0]
        if cids:
            cite_runs(pp, cids, size=8.5)
        else:
            _set_run_font(pp.add_run("—"), size=9.3)

    # 11. 관련 법령 및 근거
    H(10)
    for cv in cviews.values():
        bp = doc.add_paragraph(style="List Bullet")
        _add_hyperlink(bp, cv.href, cv.label, size=10.5)
        if cv.quote:
            qp = doc.add_paragraph()
            _set_run_font(qp.add_run("　“" + cv.quote[:150].strip() + "”"), size=9,
                          rgb=(0x40, 0x40, 0x40))

    # 12. 회계사 검토 필요사항
    H(11)
    for rp in plan.review_points:
        bp = doc.add_paragraph(style="List Bullet"); _set_run_font(bp.add_run(_ko(rp)), size=10.5)

    # 13. 자료·가정의 한계
    H(12)
    for dl in plan.data_limits:
        bp = doc.add_paragraph(style="List Bullet")
        _set_run_font(bp.add_run(_ko(dl)), size=9.8, rgb=(0xB0, 0x60, 0x00))

    # 14. 결론
    H(13); body(plan.conclusion)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
