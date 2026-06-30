"""src/tax_plan_report.py — 시나리오 기반 Tax Plan 회계사용 DOCX 렌더러.

``TaxPlanData`` 를 (1)재무제표(쟁점 행 강조표시+라벨+근거 링크), (2)시나리오 비교표,
(3)그림·도표(지분구조도·타임라인·세부담 비교·증여세 워터폴), (4)근거 법령 하이퍼링크,
(5)Noto Sans KR·영어 0 으로 렌더한다. 표현 헬퍼(_ko·하이퍼링크·폰트·음영)는 ``src.cpa_report``
의 것을 재사용한다. 인용은 등록 소스객체로 검증(날조 차단)한다.
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from src.cpa_report import (
    _FONT, _HEADING_RGB, _add_hyperlink, _apply_document_style, _ko, _qn,
    _set_run_font, _shade,
)
from src.draft import CitationView, _law_permalink
from src.source_registry import CitationVerificationError, SourceRegistry
from src.tax_plan import TaxPlanData

_ISSUE_FILL = "FEF7E0"      # 쟁점 행 강조(연한 노랑)
_SUBHEAD_FILL = "F1F3F4"    # 구분행(연한 회색)
_HEADER_FILL = "E8F0FE"     # 표 머리행(연한 블루)
_GREY = (0x5F, 0x63, 0x68)

_SECTIONS = [
    "1. 핵심 요약 및 목표",
    "2. 그룹 및 지분구조",
    "3. 재무상태표 (쟁점 강조)",
    "4. 손익계산서 (쟁점 강조)",
    "5. 핵심 세무 쟁점",
    "6. 잉여금 환원 시나리오 비교",
    "7. 권고안",
    "8. 실행 타임라인 (시점 최적화)",
    "9. 가업승계 증여세",
    "10. 관련 법령 및 근거 (클릭하면 원문으로 이동)",
    "11. 회계사 검토 필요사항",
    "12. 실행 절차 및 유의사항",
    "13. 결론",
]


class TaxPlanValidationError(RuntimeError):
    """인용 날조/미해소 — 생성 실패."""


def _validate(plan: TaxPlanData) -> None:
    """렌더러 경계의 fail-closed 검증.

    (A) §10에 렌더되는 ``plan.citations`` 전부를 등록 소스객체로 검증한다(미사용·미검증
        인용도 차단). (B) 본문(재무제표 쟁점행·쟁점·시나리오·타임라인)이 참조하는 모든
        citation_id 가 그 검증된 인용으로 해소되는지 확인한다. 따라서 DOCX 어디에도
        검증 안 된 인용이 표시되지 않는다.
    """
    cidx = plan.citation_index
    registry = SourceRegistry() if plan.source_objects else None
    if registry is not None:
        registry.register_all(plan.source_objects)

    if not plan.citations:
        raise TaxPlanValidationError("근거 법령(Citation) 없음")

    # (A) 렌더되는 모든 인용 자체를 검증 — §10 관련 법령 섹션은 plan.citations 전부를 표시한다.
    for cit in plan.citations:
        if registry is None:
            raise TaxPlanValidationError(
                f"인용 미검증(날조 차단): {cit.citation_id} — source_objects 필요")
        try:
            registry.require_citation(cit)
        except CitationVerificationError as exc:
            raise TaxPlanValidationError(f"인용 검증 실패: {cit.citation_id} — {exc}") from exc

    # (B) 본문에서 참조하는 citation_id 가 전부 (검증된) 인용으로 해소되는지 확인.
    def check(ids: list[str], where: str) -> None:
        for cid in ids:
            if cidx.get(cid) is None:
                raise TaxPlanValidationError(f"인용 미해소: {where} 의 {cid}")

    for ln in (*plan.balance_sheet.lines, *plan.income_statement.lines):
        if ln.is_issue:
            check(ln.citation_ids, f"재무제표 쟁점행 '{ln.label}'")
    for it in plan.issues:
        check(it.citation_ids, f"쟁점 '{it.title}'")
    for sc in plan.scenarios:
        check(sc.citation_ids, f"시나리오 '{sc.label}'")
    for st in plan.timeline:
        check(st.citation_ids, f"타임라인 '{st.action}'")
    if not any(s.recommended for s in plan.scenarios):
        raise TaxPlanValidationError("권고 시나리오(recommended) 1건 필요")


def _font_cell(cell, text, *, size=9.5, bold=False, rgb=None, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(_ko(text))
    _set_run_font(run, size=size, bold=bold, rgb=rgb)
    if fill:
        _shade(cell, fill)
    return p


def _issue_cell(cell, label, cviews, citation_ids, *, fill=_ISSUE_FILL):
    """쟁점 셀 — ★ + 라벨 + 근거 하이퍼링크."""
    cell.text = ""
    p = cell.paragraphs[0]
    _set_run_font(p.add_run("★ " + _ko(label) + " "), size=9, bold=True, rgb=(0xB0, 0x60, 0x00))
    for cid in citation_ids:
        cv = cviews.get(cid)
        if cv:
            _add_hyperlink(p, cv.href, cv.label, size=8.5)
            _set_run_font(p.add_run(" "), size=8.5)
    _shade(cell, fill)


def build_tax_plan_docx(
    plan: TaxPlanData,
    out_path: str | Path,
    *,
    titles: Optional[dict] = None,
    articles: Optional[dict] = None,
    charts_dir: Optional[str | Path] = None,
    with_charts: bool = True,
) -> Path:
    """시나리오 Tax Plan DOCX 생성(쟁점 강조·도표·하이퍼링크·Noto Sans KR·영어 0)."""
    from datetime import datetime

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    _validate(plan)
    cviews = {cv.citation_id: cv for cv in
              [CitationView.from_citation(c, title=(titles or {}).get(c.citation_id, ""),
                                          article_label=(articles or {}).get(c.citation_id, ""))
               for c in plan.citations]}

    out_path = Path(out_path)
    charts_dir = Path(charts_dir) if charts_dir else out_path.parent / "_charts"
    chart_paths: dict[str, Path] = {}
    if with_charts:
        from src import tax_charts as tc
        chart_paths["own"] = tc.save_ownership_diagram(
            plan.holdings_before, plan.holdings_after, charts_dir / "own.png")
        chart_paths["timeline"] = tc.save_timeline(plan.timeline, charts_dir / "timeline.png")
        chart_paths["bar"] = tc.save_scenario_bar(plan.scenarios, charts_dir / "bar.png")
        chart_paths["waterfall"] = tc.save_gift_waterfall(
            plan.gift_tax_waterfall, plan.gift_tax_before_eok, plan.gift_tax_after_eok,
            charts_dir / "waterfall.png")

    doc = Document()
    doc.core_properties.created = datetime(2024, 1, 1)
    doc.core_properties.modified = datetime(2024, 1, 1)
    _apply_document_style(doc)

    # 표지
    t = doc.add_heading(_ko(plan.title), level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(sub.add_run(f"검토기준일 {plan.as_of:%Y년 %m월 %d일}　·　{plan.group_name}　·　내부 검토용"),
                  size=10, rgb=_GREY)
    note = doc.add_paragraph(); note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(note.add_run("모든 법적 판단에는 근거 법령을 함께 표기하며, 파란색 밑줄 글씨를 누르면 "
                  "국가법령정보센터 원문으로 이동합니다. 최종 판단·서명은 담당 회계사가 합니다."),
                  size=9.5, rgb=_GREY)

    def H(i): doc.add_heading(_SECTIONS[i], level=1)
    def body(txt, size=10.5):
        p = doc.add_paragraph(); _set_run_font(p.add_run(_ko(txt)), size=size); return p

    def add_chart(key, width=6.4):
        if with_charts and key in chart_paths:
            doc.add_picture(str(chart_paths[key]), width=Inches(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. 핵심 요약 및 목표
    H(0); body(plan.executive_summary)
    p = doc.add_paragraph(); _set_run_font(p.add_run("[플랜 목표] "), size=10.5, bold=True,
                                           rgb=_HEADING_RGB)
    _set_run_font(p.add_run(_ko(plan.objective)), size=10.5)

    # 2. 그룹 및 지분구조 (+ 지분구조도)
    H(1)
    for e in plan.entities:
        bp = doc.add_paragraph(style="List Bullet"); _set_run_font(bp.add_run(_ko(e)), size=10.5)
    add_chart("own", 6.6)

    # 3. 재무상태표 / 4. 손익계산서 (쟁점 강조)
    for idx, fs in ((2, plan.balance_sheet), (3, plan.income_statement)):
        H(idx)
        cap = doc.add_paragraph()
        _set_run_font(cap.add_run(f"{fs.period}　({fs.unit})"), size=9.5, rgb=_GREY)
        tbl = doc.add_table(rows=1, cols=3); tbl.style = "Table Grid"
        for i, htxt in enumerate(["계정", "금액(억원)", "쟁점"]):
            _font_cell(tbl.rows[0].cells[i], htxt, bold=True, rgb=_HEADING_RGB,
                       fill=_HEADER_FILL, size=10)
        for ln in fs.lines:
            r = tbl.add_row().cells
            if ln.amount_eok is None:  # 구분행
                _font_cell(r[0], ln.label, bold=True, fill=_SUBHEAD_FILL, size=10)
                _font_cell(r[1], "", fill=_SUBHEAD_FILL); _font_cell(r[2], "", fill=_SUBHEAD_FILL)
                continue
            name = ("　" * ln.indent) + ln.label
            fill = _ISSUE_FILL if ln.is_issue else None
            _font_cell(r[0], name, bold=ln.bold, fill=fill, size=9.7)
            _font_cell(r[1], f"{ln.amount_eok:,.0f}", bold=ln.bold, fill=fill, size=9.7)
            if ln.is_issue:
                _issue_cell(r[2], ln.issue_label, cviews, ln.citation_ids)
            else:
                _font_cell(r[2], "", fill=fill)

    # 5. 핵심 세무 쟁점
    H(4)
    for it in plan.issues:
        sev = {"HIGH": "높음", "MEDIUM": "중간", "LOW": "낮음"}.get(it.severity, "중간")
        bp = doc.add_paragraph(style="List Bullet")
        _set_run_font(bp.add_run(f"[중요도 {sev}] {_ko(it.title)} — "), size=10.5, bold=True)
        _set_run_font(bp.add_run(_ko(it.description)), size=10.5)
        if it.citation_ids:
            _set_run_font(bp.add_run("　(근거: "), size=9, rgb=_GREY)
            for j, cid in enumerate([c for c in it.citation_ids if c in cviews]):
                if j:
                    _set_run_font(bp.add_run(", "), size=9)
                _add_hyperlink(bp, cviews[cid].href, cviews[cid].label, size=9)
            _set_run_font(bp.add_run(")"), size=9, rgb=_GREY)

    # 6. 시나리오 비교 (표 + 막대차트)
    H(5)
    body("잉여금 환원 방법(현금배당·유상감자·혼합)별로 추가 세부담과 승계가치 영향을 비교합니다.", 9.5)
    st = doc.add_table(rows=1, cols=4); st.style = "Table Grid"
    for i, htxt in enumerate(["시나리오", "요지", "추가 세부담", "장단점"]):
        _font_cell(st.rows[0].cells[i], htxt, bold=True, rgb=_HEADING_RGB, fill=_HEADER_FILL, size=10)
    for sc in plan.scenarios:
        r = st.add_row().cells
        fill = "E6F4EA" if sc.recommended else None
        lab = f"{sc.key}. {sc.label}" + ("　◀ 권고" if sc.recommended else "")
        _font_cell(r[0], lab, bold=sc.recommended, fill=fill, size=9.5)
        _font_cell(r[1], sc.summary, fill=fill, size=9)
        _font_cell(r[2], f"약 {sc.tax_burden_eok:,.0f}억원", bold=sc.recommended, fill=fill, size=9.5)
        _font_cell(r[3], f"장점: {sc.pros}\n단점: {sc.cons}", fill=fill, size=9)
    # 각 시나리오 세액계산 표
    for sc in plan.scenarios:
        h2 = doc.add_heading(f"시나리오 {sc.key}. {_ko(sc.label)}", level=2)
        ct = doc.add_table(rows=1, cols=3); ct.style = "Table Grid"
        for i, htxt in enumerate(["계산 항목", "금액", "비고"]):
            _font_cell(ct.rows[0].cells[i], htxt, bold=True, fill=_HEADER_FILL, size=9.5)
        for row in sc.rows:
            r = ct.add_row().cells
            _font_cell(r[0], row.item, size=9.3); _font_cell(r[1], row.amount, size=9.3)
            _font_cell(r[2], row.note, size=9.3)
    add_chart("bar", 6.2)

    # 7. 권고안
    H(6); body(plan.recommended_note)

    # 8. 실행 타임라인 (+ 도식)
    H(7)
    add_chart("timeline", 6.8)
    tt = doc.add_table(rows=1, cols=4); tt.style = "Table Grid"
    for i, htxt in enumerate(["시점", "행위", "세무 효과", "근거"]):
        _font_cell(tt.rows[0].cells[i], htxt, bold=True, fill=_HEADER_FILL, size=9.5)
    for s in plan.timeline:
        r = tt.add_row().cells
        _font_cell(r[0], s.date_label, bold=True, size=9.3); _font_cell(r[1], s.action, size=9.3)
        _font_cell(r[2], s.tax_effect, size=9.3)
        if s.citation_ids:
            cell = r[3]; cell.text = ""
            pp = cell.paragraphs[0]
            for j, cid in enumerate([c for c in s.citation_ids if c in cviews]):
                if j:
                    _set_run_font(pp.add_run(" "), size=8.5)
                _add_hyperlink(pp, cviews[cid].href, cviews[cid].label, size=8.5)
        else:
            _font_cell(r[3], "—", size=9.3)

    # 9. 가업승계 증여세 (+ 워터폴)
    H(8); body(plan.gift_tax_note)
    add_chart("waterfall", 6.4)
    p = doc.add_paragraph()
    _set_run_font(p.add_run(f"정리 전 증여세 약 {plan.gift_tax_before_eok:,.0f}억원 → 정리 후 약 "
                  f"{plan.gift_tax_after_eok:,.0f}억원 (약 "
                  f"{plan.gift_tax_before_eok - plan.gift_tax_after_eok:,.0f}억원 절감, 추정)"),
                  size=10.5, bold=True, rgb=_HEADING_RGB)

    # 10. 관련 법령 및 근거 (하이퍼링크)
    H(9)
    for cv in cviews.values():
        bp = doc.add_paragraph(style="List Bullet")
        _add_hyperlink(bp, cv.href, cv.label, size=10.5)
        if cv.quote:
            qp = doc.add_paragraph()
            _set_run_font(qp.add_run("　“" + cv.quote[:150].strip() + "”"), size=9,
                          rgb=(0x40, 0x40, 0x40))

    # 11. 회계사 검토 필요사항
    H(10)
    for rp in plan.review_points:
        bp = doc.add_paragraph(style="List Bullet"); _set_run_font(bp.add_run(_ko(rp)), size=10.5)

    # 12. 실행 절차 및 유의사항
    H(11)
    for pn in plan.procedure_notes:
        bp = doc.add_paragraph(style="List Bullet"); _set_run_font(bp.add_run(_ko(pn)), size=10.5)
    if plan.data_limits:
        dp = doc.add_paragraph()
        _set_run_font(dp.add_run("자료·가정 한계: " + " / ".join(_ko(x) for x in plan.data_limits)),
                      size=9, rgb=(0xB0, 0x60, 0x00))

    # 13. 결론
    H(12); body(plan.conclusion)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
